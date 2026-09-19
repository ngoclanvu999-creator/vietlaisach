# -*- coding: utf-8 -*-
"""
Dựng tệp Word cho app SOẠN CHUYÊN ĐỀ.

Chỉ hai loại, nên KHÔNG có bảng rẽ nhánh chín nhánh như công cụ cũ. Chính bảng
rẽ nhánh đó là thứ làm mọi thứ rối: sửa chuyên đề phải mở tệp mà giáo án, đề
thi và bài giảng cùng nằm trong đó.
"""

from pathlib import Path
from typing import Any, Dict

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from core.exporter_nen import COLOR_TEXT_MUTED, FONT_MAIN, NenWord
from core.loai_dau_ra import CHUYEN_DE_BT, TAI_LIEU_HSG
from core.trang_tri import cach_bay

from .exporter_chuyen_de import xuat_chuyen_de
from .exporter_hsg import xuat_tai_lieu_hsg


def _dau_trang(doc, tieu_de: str):
    """Đầu trang và chân trang theo thể thức tài liệu nội bộ."""
    tiet = doc.sections[0]
    p = tiet.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"TÀI LIỆU LƯU HÀNH NỘI BỘ — CHƯƠNG TRÌNH GDPT 2018 | {tieu_de}")
    r.font.name = FONT_MAIN
    r.font.size = Pt(9)
    r.font.color.rgb = COLOR_TEXT_MUTED

    p2 = tiet.footer.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Biên soạn tự động — chuẩn thể thức Nghị định 30/2020/NĐ-CP")
    r2.font.name = FONT_MAIN
    r2.font.size = Pt(9)
    r2.font.color.rgb = COLOR_TEXT_MUTED


def dung_tai_lieu(book, duong_ra: Path, kho_giay: str = "a4",
                  tuy_chon: Dict[str, Any] = None) -> Path:
    """
    Dựng trọn tệp .docx rồi lưu.

    `book.loai_dau_ra` quyết định gọi bộ dựng nào — chỉ hai nhánh, và cả hai
    đều nằm ngay trong thư mục app này.
    """
    opts = NenWord.merge_options(tuy_chon)
    loai = getattr(book, "loai_dau_ra", "") or CHUYEN_DE_BT
    mon = getattr(book, "subject", "") or "toan"
    cb = cach_bay(getattr(book, "cap_hoc", "THPT"), loai)

    doc = docx.Document()
    NenWord.apply_default_style(doc, cb)
    NenWord.apply_standard_page_setup(doc.sections[0], paper_format=kho_giay)
    _dau_trang(doc, getattr(book, "new_title", "") or "")

    if loai == TAI_LIEU_HSG:
        xuat_tai_lieu_hsg(doc, book, mon, opts)
    else:
        xuat_chuyen_de(doc, book, mon, opts)

    duong_ra.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(duong_ra))
    return duong_ra
