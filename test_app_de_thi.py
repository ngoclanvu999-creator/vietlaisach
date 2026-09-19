# -*- coding: utf-8 -*-
"""
Kiểm thử app SOẠN ĐỀ THI — dựng file Word thật rồi mở lại đo.

Trước đây KHÔNG hề có bộ kiểm thử nào cho năm loại đề, dù đây là nhóm nghiệp vụ
nặng nhất: ma trận, bản đặc tả, ba phần I/II/III, hướng dẫn chấm, và thang điểm
phải cộng đúng 10,0 ở mọi tổ hợp loại đề × môn.

Chỗ dễ lập trình sai nhất là **thang điểm Phần II**: lũy tiến
0,1 / 0,25 / 0,5 / 1,0 theo số ý đúng, KHÔNG phải 0,25 nhân số ý. Đúng 2 trong
4 ý chỉ được 0,25 — nếu nhân đều thì thành 0,5, sai gấp đôi.
"""

import sys
from pathlib import Path

import docx

from app_de_thi.bo_dung.dung import dung_tai_lieu
from app_de_thi.cau_hinh import LOAI_DAU_RA
from core.loai_dau_ra import (DE_CHUONG_BAI, DE_GIUA_KI_1, DE_GIUA_KI_2,
                              DE_HOC_KI_1, DE_HOC_KI_2, QUY_CACH,
                              kiem_tra_tong_diem, tinh_cau_truc)
from core.rewriter import RewrittenBook, RewrittenQuestionItem

RA = Path("app_de_thi/output")
MUC = ["Nhận biết", "Thông hiểu", "Vận dụng"]


class Soat:
    def __init__(self, ten):
        self.loi = 0
        print("\n" + ten)
        print("-" * 78)

    def __call__(self, ten: str, ok: bool, them: str = ""):
        print("  %-5s %s" % ("ĐẠT" if ok else "HỎNG", ten))
        if them:
            print("        " + them)
        if not ok:
            self.loi += 1


def dung_sach(loai: str, mon: str, so_cau: int) -> RewrittenBook:
    ds = []
    for i in range(1, so_cau + 1):
        q = RewrittenQuestionItem(index=i)
        q.new_content = f"Câu hỏi thử số {i} về nội dung trọng tâm chương trình."
        q.level = MUC[i % 3]
        q.correct_answer = "ABCD"[i % 4]
        q.new_options = ["A. 1", "B. 2", "C. 3", "D. 4"]
        q.solution_method1 = f"Lời giải mẫu cho câu {i}."
        ds.append(q)
    return RewrittenBook(
        original_title="Thu", new_title=f"Đề thử {QUY_CACH[loai].ten}",
        subtitle="", author_note="", chapter_summary="",
        questions=ds, loai_dau_ra=loai, cap_hoc="THPT", subject=mon,
    )


def kiem_danh_muc() -> int:
    """App phải nhận ĐÚNG năm loại đề, không thừa không thiếu."""
    k = Soat("DANH MỤC LOẠI ĐỀ")
    mong_doi = [DE_CHUONG_BAI, DE_GIUA_KI_1, DE_GIUA_KI_2, DE_HOC_KI_1, DE_HOC_KI_2]
    k("Đúng năm loại đề", list(LOAI_DAU_RA) == mong_doi,
      " · ".join(QUY_CACH[m].ten for m in LOAI_DAU_RA))
    k("Không lẫn loại của app khác",
      all(m.startswith("DE_") for m in LOAI_DAU_RA))
    return k.loi


def kiem_thang_diem() -> int:
    """
    Điều quan trọng nhất, và kiểm bằng TÍNH LẠI ĐỘC LẬP chứ không gọi hàm của
    chính mã nguồn — gọi lại hàm của nó thì nó sai kiểu gì cũng vẫn "đạt".
    """
    k = Soat("THANG ĐIỂM — mọi tổ hợp phải cộng đúng 10,0")
    for loai in LOAI_DAU_RA:
        for mon in ("toan", "vatly"):
            ct = tinh_cau_truc(loai, mon)
            tong = (ct.so_cau_p1 * ct.diem_moi_cau_p1
                    + ct.so_cau_p2 * ct.diem_toi_da_p2
                    + ct.so_cau_p3 * ct.diem_moi_cau_p3)
            k("%-14s %-6s = %.2f điểm" % (loai, mon, tong), abs(tong - 10.0) < 0.001,
              "%d + %d + %d = %d câu" % (ct.so_cau_p1, ct.so_cau_p2, ct.so_cau_p3,
                                         ct.tong_cau))
            k("   hàm kiem_tra_tong_diem đồng ý", kiem_tra_tong_diem(ct))

    # Thang lũy tiến Phần II — chỗ dễ sai nhất
    ct = tinh_cau_truc(DE_HOC_KI_1, "toan")
    k("Phần II lũy tiến đúng 0,1 / 0,25 / 0,5 / 1,0",
      ct.luy_tien_p2 == [0.1, 0.25, 0.5, 1.0], str(ct.luy_tien_p2))
    k("Đúng 2/4 ý CHỈ được 0,25 (không phải 0,5)",
      ct.luy_tien_p2[1] == 0.25,
      "nếu nhân đều 0,25 × 2 ý thì thành 0,50 — sai gấp đôi")

    # Đề học kì phải khớp con số Bộ đã công bố
    ct_t = tinh_cau_truc(DE_HOC_KI_1, "toan")
    ct_l = tinh_cau_truc(DE_HOC_KI_1, "vatly")
    k("Toán học kì: 12 + 4 + 6 = 22 câu",
      (ct_t.so_cau_p1, ct_t.so_cau_p2, ct_t.so_cau_p3) == (12, 4, 6),
      "đo được %d + %d + %d" % (ct_t.so_cau_p1, ct_t.so_cau_p2, ct_t.so_cau_p3))
    k("Vật lí học kì: 18 + 4 + 6 = 28 câu",
      (ct_l.so_cau_p1, ct_l.so_cau_p2, ct_l.so_cau_p3) == (18, 4, 6),
      "đo được %d + %d + %d" % (ct_l.so_cau_p1, ct_l.so_cau_p2, ct_l.so_cau_p3))
    return k.loi


def kiem_bo_bon_phan() -> int:
    ct = tinh_cau_truc(DE_HOC_KI_1, "toan")
    duong = RA / "Đề thi học kì 1" / "Test_De_Hoc_Ki.docx"
    dung_tai_lieu(dung_sach(DE_HOC_KI_1, "toan", ct.tong_cau), duong)

    d = docx.Document(str(duong))
    dong = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    gop = "\n".join(dong)
    k = Soat("BỘ BỐN PHẦN — đo trên tệp .docx thật")

    for ten in ("MA TRẬN", "ĐẶC TẢ", "HƯỚNG DẪN CHẤM"):
        k("Có phần %s" % ten, ten in gop.upper())

    for phan in ("PHẦN I", "PHẦN II", "PHẦN III"):
        k("Đề có %s" % phan, phan in gop.upper())

    # Thứ tự: ma trận trước, hướng dẫn chấm sau cùng
    hoa = gop.upper()
    k("Ma trận đứng trước hướng dẫn chấm",
      hoa.find("MA TRẬN") < hoa.find("HƯỚNG DẪN CHẤM"))

    k("Có bảng (ma trận và đặc tả đều là bảng)", len(d.tables) >= 2,
      "%d bảng" % len(d.tables))

    print("        Tệp: %s" % duong)
    return k.loi


def kiem_the_thuc() -> int:
    duong = RA / "Đề thi học kì 1" / "Test_De_Hoc_Ki.docx"
    s = docx.Document(str(duong)).sections[0]
    k = Soat("THỂ THỨC NGHỊ ĐỊNH 30 (đo trên tệp)")
    mm = lambda emu: round(emu / 36000)
    k("Lề trái 30mm để đóng gáy", mm(s.left_margin) == 30, "đo được %dmm" % mm(s.left_margin))
    k("Lề phải 15mm", mm(s.right_margin) == 15, "đo được %dmm" % mm(s.right_margin))
    k("Khổ A4 rộng 210mm", mm(s.page_width) == 210, "đo được %dmm" % mm(s.page_width))
    return k.loi


def kiem_moi_loai_dung_duoc() -> int:
    """Cả năm loại đều phải dựng ra tệp, không loại nào ném ngoại lệ."""
    k = Soat("CẢ NĂM LOẠI ĐỀ ĐỀU DỰNG ĐƯỢC")
    for loai in LOAI_DAU_RA:
        ct = tinh_cau_truc(loai, "toan")
        duong = RA / "thu" / ("%s.docx" % loai)
        try:
            dung_tai_lieu(dung_sach(loai, "toan", ct.tong_cau), duong)
            co = duong.exists() and duong.stat().st_size > 8000
            k("%-14s (%d câu, %d phút)" % (loai, ct.tong_cau, ct.thoi_gian), co,
              "%d byte" % duong.stat().st_size if duong.exists() else "không tạo được tệp")
        except Exception as e:
            k("%-14s" % loai, False, "%s: %s" % (type(e).__name__, str(e)[:90]))
    return k.loi


if __name__ == "__main__":
    loi = (kiem_danh_muc() + kiem_thang_diem() + kiem_bo_bon_phan()
           + kiem_the_thuc() + kiem_moi_loai_dung_duoc())
    print()
    print("=" * 78)
    print("KẾT QUẢ: %s" % ("ĐẠT" if loi == 0 else "%d LỖI" % loi))
    sys.exit(1 if loi else 0)
