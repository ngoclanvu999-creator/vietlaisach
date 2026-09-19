# -*- coding: utf-8 -*-
"""
Dựng tệp Word cho app SOẠN ĐỀ THI.

Đầu ra LUÔN là bộ bốn phần theo chuẩn 2025: ma trận đề · bản đặc tả · đề ba
phần I/II/III · hướng dẫn chấm. Không có kiểu "đề gọn" nào nữa.

Bộ dựng đề CŨ (`render_exam_header` / `export_exam` trong `core/exporter.py`,
khoảng 220 dòng) đã bị xóa ngày 19/09/2026 theo quyết định của chủ dự án. Giữ
hai bộ dựng cho cùng một việc là cách chắc chắn nhất để chúng trôi khác nhau.
"""

from pathlib import Path
from typing import Any, Dict

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from core.exporter_nen import COLOR_TEXT_MUTED, FONT_MAIN, NenWord
from core.trang_tri import cach_bay

from .exporter_de import xuat_bo_de


def _dau_trang(doc, tieu_de: str):
    tiet = doc.sections[0]
    p = tiet.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"ĐỀ KIỂM TRA — CHƯƠNG TRÌNH GDPT 2018 | {tieu_de}")
    r.font.name = FONT_MAIN
    r.font.size = Pt(9)
    r.font.color.rgb = COLOR_TEXT_MUTED

    p2 = tiet.footer.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Biên soạn tự động — chuẩn thể thức Nghị định 30/2020/NĐ-CP")
    r2.font.name = FONT_MAIN
    r2.font.size = Pt(9)
    r2.font.color.rgb = COLOR_TEXT_MUTED


def dung_tai_lieu(book, duong_ra: Path, kho_giay: str = "a4",
                  tuy_chon: Dict[str, Any] = None) -> Path:
    """Dựng trọn bộ đề bốn phần rồi lưu."""
    opts = NenWord.merge_options(tuy_chon)
    loai = getattr(book, "loai_dau_ra", "") or ""
    mon = getattr(book, "subject", "") or "toan"

    # Đề kiểm tra là hồ sơ chuyên môn: mọi cấp học đều giữ nghiêm ngặt, không
    # biểu tượng, không khung màu.
    cb = cach_bay(getattr(book, "cap_hoc", "THPT"), loai)

    doc = docx.Document()
    NenWord.apply_default_style(doc, cb)
    NenWord.apply_standard_page_setup(doc.sections[0], paper_format=kho_giay)
    _dau_trang(doc, getattr(book, "new_title", "") or "")

    xuat_bo_de(doc, book, loai, mon, kem_loi_giai=bool(opts.get("solution", True)))

    duong_ra.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(duong_ra))
    return duong_ra
