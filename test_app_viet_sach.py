# -*- coding: utf-8 -*-
"""
Kiểm thử app VIẾT LẠI SÁCH — dựng file Word thật rồi mở lại đo.

App này chỉ làm một việc, nhưng đó là việc gốc của cả dự án: biến tài liệu rời
rạc thành một cuốn sách đọc được — bìa, lời tựa, các chương, lời giải.

Gộp phần kiểm của `test_pipeline.py` và `test_upgrade.py` vốn chỉ chạy luồng
chung mà không đo gì trên tệp thành phẩm.
"""

import sys
from pathlib import Path

import docx

from app_viet_sach.bo_dung.dung import dung_tai_lieu
from core.rewriter import (RewrittenBook, RewrittenChapter,
                           RewrittenQuestionItem)

RA = Path("app_viet_sach/output")


class Soat:
    def __init__(self, ten):
        self.loi = 0
        print("\n" + ten)
        print("-" * 78)

    def __call__(self, ten: str, ok: bool, them: str = ""):
        print("  %-5s %s" % ("ĐẠT" if ok else "HỎNG", ten))
        if them:
            print("        " + them)
        if not ok:
            self.loi += 1


def _cau(i: int) -> RewrittenQuestionItem:
    q = RewrittenQuestionItem(index=i)
    q.title = f"Câu {i}"
    q.new_content = f"Cho hàm số y = x^2 - {i}x + 1. Tìm tọa độ đỉnh của parabol."
    q.level = ["Nhận biết", "Thông hiểu", "Vận dụng"][i % 3]
    q.new_options = ["A. (1;0)", "B. (2;1)", "C. (0;1)", "D. (3;2)"]
    q.correct_answer = "ABCD"[i % 4]
    q.solution_method1 = (f"Đỉnh parabol có hoành độ x = {i}/2. "
                          "Thay vào hàm số được tung độ tương ứng.")
    q.solution_method2 = "Bấm MODE 5 3 rồi nhập hệ số."
    q.trap_warning = "Hay nhầm dấu khi tính -b/2a."
    return q


def dung_sach_don() -> RewrittenBook:
    return RewrittenBook(
        original_title="Tai lieu goc", new_title="Cẩm Nang Hàm Số Bậc Hai",
        subtitle="Hệ thống kiến thức và bài tập chọn lọc",
        author_note="Cuốn sách này dành cho học sinh muốn nắm chắc phần hàm số.",
        chapter_summary="Hàm số bậc hai",
        theory_section="I. ĐỊNH NGHĨA\nHàm số bậc hai có dạng y = ax^2 + bx + c.",
        questions=[_cau(i) for i in range(1, 9)],
        cap_hoc="THPT", subject="toan",
    )


def dung_sach_nhieu_chuong() -> RewrittenBook:
    ch = [
        RewrittenChapter(index=1, title="Chương 1. Hàm số bậc hai",
                         source_name="tep_a.docx", theory_section="",
                         questions=[_cau(i) for i in range(1, 5)]),
        RewrittenChapter(index=2, title="Chương 2. Phương trình bậc hai",
                         source_name="tep_b.docx", theory_section="",
                         questions=[_cau(i) for i in range(5, 9)]),
    ]
    return RewrittenBook(
        original_title="Gop thu muc", new_title="Tuyển Tập Đại Số Lớp 10",
        subtitle="Gộp từ nhiều tài liệu", author_note="Lời tựa mẫu.",
        chapter_summary="", theory_section="", chapters=ch,
        questions=[q for c in ch for q in c.questions],
        cap_hoc="THPT", subject="toan",
    )


def toan_bo_chu(d) -> str:
    """
    Gom chữ của CẢ đoạn văn LẪN bảng.

    Bộ dựng sách đặt lời tựa, lý thuyết nền và mẹo Casio trong hộp ghi chú — mà
    hộp ghi chú là một bảng. Chỉ đọc `doc.paragraphs` thì không thấy chúng và
    phép kiểm sẽ báo thiếu oan.
    """
    phan = [x.text for x in d.paragraphs]
    for b in d.tables:
        for r in b.rows:
            for o in r.cells:
                phan.append(o.text)
    return "\n".join(t.strip() for t in phan if t.strip())


def kiem_sach_don() -> int:
    duong = RA / "Test_Sach_Don.docx"
    dung_tai_lieu(dung_sach_don(), duong, tuy_chon={"solution": True, "casio": True,
                                                    "traps": True, "theory": True})
    d = docx.Document(str(duong))
    dong = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    gop = toan_bo_chu(d)
    k = Soat("SÁCH MỘT CUỐN — đo trên tệp .docx thật")

    k("Có tên sách trên trang bìa", "CẨM NANG HÀM SỐ BẬC HAI" in gop.upper())
    k("Có lời tựa của tác giả", "dành cho học sinh muốn nắm chắc" in gop)
    k("Có phần lý thuyết nền", "ĐỊNH NGHĨA" in gop.upper())
    so_cau = sum(1 for t in dong if t.lstrip("▶ ").startswith("Câu "))
    k("Đủ tám câu hỏi", so_cau >= 8, "đếm được %d câu" % so_cau)
    k("Có lời giải", "Đỉnh parabol" in gop)
    k("Có cảnh báo bẫy", "nhầm dấu" in gop)
    k("Có mẹo Casio khi bật tùy chọn", "MODE 5 3" in gop)
    print("        Tệp: %s (%d đoạn, %d bảng)" % (duong, len(dong), len(d.tables)))
    return k.loi


def kiem_sach_nhieu_chuong() -> int:
    duong = RA / "Test_Sach_Nhieu_Chuong.docx"
    dung_tai_lieu(dung_sach_nhieu_chuong(), duong)
    gop = toan_bo_chu(docx.Document(str(duong)))
    k = Soat("SÁCH NHIỀU CHƯƠNG")
    k("Có tên chương 1", "Chương 1" in gop)
    k("Có tên chương 2", "Chương 2" in gop)
    k("Chương 1 đứng trước chương 2", gop.find("Chương 1") < gop.find("Chương 2"))
    print("        Tệp: %s" % duong)
    return k.loi


def kiem_the_thuc() -> int:
    """Lề Nghị định 30 đo trên TỆP THẬT — bộ thẩm định từng chấm cứng 20/20 mà không mở file."""
    k = Soat("THỂ THỨC NGHỊ ĐỊNH 30 (đo trên tệp)")
    for ten, kho, rong_mm in (("Test_Sach_Don.docx", "a4", 210),):
        s = docx.Document(str(RA / ten)).sections[0]
        mm = lambda emu: round(emu / 36000)
        k("Lề trái 30mm để đóng gáy", mm(s.left_margin) == 30,
          "đo được %dmm" % mm(s.left_margin))
        k("Lề phải 15mm", mm(s.right_margin) == 15, "đo được %dmm" % mm(s.right_margin))
        k("Khổ %s rộng %dmm" % (kho.upper(), rong_mm), mm(s.page_width) == rong_mm,
          "đo được %dmm" % mm(s.page_width))

    # Khổ B5 phải khác A4 thật sự, không phải chỉ đổi nhãn
    duong_b5 = RA / "Test_Sach_B5.docx"
    dung_tai_lieu(dung_sach_don(), duong_b5, kho_giay="b5")
    s5 = docx.Document(str(duong_b5)).sections[0]
    k("Khổ B5 rộng 176mm", round(s5.page_width / 36000) == 176,
      "đo được %dmm" % round(s5.page_width / 36000))
    return k.loi


if __name__ == "__main__":
    loi = kiem_sach_don() + kiem_sach_nhieu_chuong() + kiem_the_thuc()
    print()
    print("=" * 78)
    print("KẾT QUẢ: %s" % ("ĐẠT" if loi == 0 else "%d LỖI" % loi))
    sys.exit(1 if loi else 0)
