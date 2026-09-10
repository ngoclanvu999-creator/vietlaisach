import sys
from pathlib import Path
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass
import docx
from docx.shared import Pt, Inches
import openpyxl

BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"
INPUT_DIR.mkdir(exist_ok=True)

def create_sample_docx():
    doc = docx.Document()
    
    # Title
    p_title = doc.add_paragraph()
    r = p_title.add_run("BỘ ĐỀ ÔN TẬP TRỌNG ĐIỂM MÔN TOÁN 12 - NĂM HỌC 2026")
    r.bold = True
    r.font.size = Pt(16)

    # Question 1
    doc.add_paragraph("Câu 1: Cho hàm số y = f(x) có đạo hàm f'(x) = (x - 1)(x + 2)²(x - 3). Số điểm cực trị của hàm số đã cho là:")
    doc.add_paragraph("A. 1\nB. 2\nC. 3\nD. 4")
    doc.add_paragraph("Lời giải: Đạo hàm f'(x) đổi dấu khi qua các nghiệm đơn x = 1 và x = 3. Nghiệm x = -2 là nghiệm bội chẵn nên đạo hàm không đổi dấu. Do đó hàm số có 2 điểm cực trị. Đáp án B.")

    # Question 2
    doc.add_paragraph("Câu 2: Tập nghiệm của bất phương trình log₂(x - 1) ≤ 3 là:")
    doc.add_paragraph("A. (1; 9]\nB. (-∞; 9]\nC. [1; 9]\nD. (1; 8]")
    doc.add_paragraph("Lời giải: Điều kiện xác định: x - 1 > 0 ⇔ x > 1. Bất phương trình tương đương: x - 1 ≤ 2³ = 8 ⇔ x ≤ 9. Kết hợp điều kiện ta có S = (1; 9]. Đáp án A.")

    # Question 3
    doc.add_paragraph("Câu 3: Tính tích phân I = ∫[0 đến 1] (2x + 1)e^x dx:")
    doc.add_paragraph("A. I = e + 1\nB. I = 2e - 1\nC. I = e - 1\nD. I = 2e + 1")
    doc.add_paragraph("Lời giải: Sử dụng phương pháp tích phân từng phần: đặt u = 2x + 1 ⇒ du = 2dx; dv = e^x dx ⇒ v = e^x. Ta có I = (2x + 1)e^x |[0,1] - ∫[0,1] 2e^x dx = 3e - 1 - 2(e - 1) = e + 1. Đáp án A.")

    # Question 4
    doc.add_paragraph("Câu 4: Cho hình chóp S.ABCD có đáy ABCD là hình vuông cạnh a, SA vuông góc với mặt phẳng (ABCD) và SA = a√3. Thể tích khối chóp S.ABCD là:")
    doc.add_paragraph("A. a³√3 / 3\nB. a³√3\nC. a³ / 3\nD. a³√3 / 6")
    doc.add_paragraph("Lời giải: Diện tích đáy S_ABCD = a². Chiều cao h = SA = a√3. Thể tích V = 1/3 * S_ABCD * h = 1/3 * a² * a√3 = a³√3 / 3. Đáp án A.")

    # Question 5
    doc.add_paragraph("Câu 5: Có bao nhiêu giá trị nguyên của tham số m để phương trình 4^x - 2^(x+1) + m = 0 có hai nghiệm phân biệt?")
    doc.add_paragraph("A. 0\nB. 1\nC. 2\nD. Vô số")
    doc.add_paragraph("Lời giải: Đặt t = 2^x (t > 0). Phương trình trở thành t² - 2t + m = 0 ⇔ m = -t² + 2t. Để phương trình có 2 nghiệm phân biệt thì đường thẳng y = m phải cắt parabol y = -t² + 2t tại 2 điểm có hoành độ dương t > 0. Khảo sát hàm g(t) = -t² + 2t với t > 0 ta có đỉnh (1; 1) và g(0) = 0. Do đó 0 < m < 1. Không có số nguyên nào thỏa mãn. Đáp án A.")

    file_path = INPUT_DIR / "sample_toan_goc.docx"
    doc.save(str(file_path))
    print(f"Created: {file_path}")

def create_sample_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NganHangCauHoi"

    # Header
    ws.append(["STT", "Đề bài", "Phương án A", "Phương án B", "Phương án C", "Phương án D", "Đáp án đúng", "Lời giải chi tiết"])

    # Data
    rows = [
        [1, "Tìm tập xác định của hàm số y = (x - 2)^(-3):", "D = R", "D = R \\ {2}", "D = (2; +∞)", "D = [2; +∞)", "B", "Số mũ nguyên âm nên điều kiện là cơ số khác 0: x - 2 ≠ 0 ⇔ x ≠ 2."],
        [2, "Cho khối lăng trụ tam giác đều có tất cả các cạnh bằng a. Thể tích của khối lăng trụ đó là:", "a³√3 / 4", "a³√3 / 12", "a³ / 4", "a³√3 / 6", "A", "Diện tích tam giác đều cạnh a là S = a²√3 / 4. Chiều cao h = a. Thể tích V = S * h = a³√3 / 4."],
        [3, "Nguyên hàm của hàm số f(x) = cos 2x là:", "1/2 sin 2x + C", "-1/2 sin 2x + C", "2 sin 2x + C", "-2 sin 2x + C", "A", "Áp dụng công thức nguyên hàm cơ bản: ∫ cos(ax + b)dx = 1/a sin(ax + b) + C với a = 2."],
        [4, "Số giao điểm của đồ thị hàm số y = x³ - 3x và trục hoành là:", "1", "2", "3", "0", "C", "Phương trình hoành độ giao điểm: x³ - 3x = 0 ⇔ x(x² - 3) = 0 có 3 nghiệm x = 0, x = √3, x = -√3."]
    ]

    for r in rows:
        ws.append(r)

    file_path = INPUT_DIR / "sample_toan_goc.xlsx"
    wb.save(str(file_path))
    print(f"Created: {file_path}")

if __name__ == "__main__":
    create_sample_docx()
    create_sample_xlsx()
