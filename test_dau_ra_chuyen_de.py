# -*- coding: utf-8 -*-
"""
Kiểm thử loại đầu ra CHUYEN_DE_BT — dựng file Word thật rồi mở lại đo.

Không tin giá trị hàm trả về: hàm báo "đã in 4 phần" mà file chỉ có 3 thì báo
cáo đó chỉ là trang trí. Mọi khẳng định dưới đây đều đếm trên tệp .docx.

Điều quan trọng nhất phải giữ: **đáp án KHÔNG được xuất hiện trước Phần 4**.
Skill chuyen-de-bai-tap đo từ mẫu thật ghi rõ "học sinh phải tự làm trước rồi
mới đối chiếu" — in đáp án ngay dưới đề là hỏng hẳn giá trị sư phạm.
"""

import sys
from pathlib import Path

import docx

from core.exporter import DocxBookExporter
from core.loai_dau_ra import CHUYEN_DE_BT
from core.rewriter import RewrittenBook, RewrittenQuestionItem

RA = Path("output") / "Test_Chuyen_De_Bon_Phan.docx"

# Câu 1-3 dễ, 4-5 trung bình, 6 khó, 7-8 có nguồn đề thi thật
CAU = [
    ("Tính giá trị của 2^3 + 3^2.", "Nhận biết", "A", "2^3 = 8, 3^2 = 9, tổng bằng 17.", ""),
    ("Tìm tập xác định của hàm số y = log(x - 1).", "Nhận biết", "B", "Điều kiện x - 1 > 0 nên x > 1.", ""),
    ("Rút gọn biểu thức 5^2 · 5^3.", "Nhận biết", "C", "Nhân cùng cơ số thì cộng số mũ: 5^5.", ""),
    ("Giải phương trình 2^x = 16.", "Thông hiểu", "A", "16 = 2^4 nên x = 4.", ""),
    ("Giải bất phương trình log2(x) < 3.", "Thông hiểu", "D", "0 < x < 8.", ""),
    ("Tìm m để phương trình 4^x - m·2^x + 1 = 0 có hai nghiệm phân biệt.",
     "Vận dụng cao", "B", "Đặt t = 2^x > 0, phương trình thành t^2 - mt + 1 = 0. "
     "Cần hai nghiệm dương phân biệt nên delta > 0 và tổng, tích dương, suy ra m > 2.", ""),
    ("Tập xác định của hàm số y = log2(3 - x) là gì?", "Thông hiểu", "A",
     "Điều kiện 3 - x > 0 nên x < 3.", "Đề thi tốt nghiệp THPT 2015"),
    ("Cho log2(a) = 3. Tính log2(4a).", "Thông hiểu", "C",
     "log2(4a) = log2(4) + log2(a) = 2 + 3 = 5.", "Đề thi tốt nghiệp THPT 2021"),
]


def dung_sach():
    ds = []
    for i, (nd, muc, dap, lg, nguon) in enumerate(CAU, 1):
        q = RewrittenQuestionItem(index=i)
        q.new_content = nd
        q.level = muc
        q.correct_answer = dap
        q.solution_method1 = lg
        q.title = nguon
        q.new_options = ["A. 17", "B. 4", "C. 5^5", "D. 8"]
        ds.append(q)

    return RewrittenBook(
        original_title="Mu va logarit", new_title="Hàm số mũ và logarit",
        subtitle="", author_note="", chapter_summary="",
        theory_section="Định nghĩa hàm số mũ\nHàm số mũ có dạng y = a^x với a > 0, a khác 1.",
        questions=ds,
        # Đi qua ĐÚNG mã loại đầu ra, không gọi tắt bộ dựng: như vậy kiểm thử
        # chứng minh được cả nhánh rẽ lẫn nội dung. Trước đây chính nhánh rẽ này
        # thiếu, làm chuyên đề rơi vào bộ dựng sách mặc định (luật KL4).
        loai_dau_ra=CHUYEN_DE_BT,
        cap_hoc="THPT",
    )


def main() -> int:
    RA.parent.mkdir(parents=True, exist_ok=True)
    DocxBookExporter.export(dung_sach(), RA)
    print("Đã dựng qua DocxBookExporter.export() với loai_dau_ra =", CHUYEN_DE_BT)
    print()

    d = docx.Document(str(RA))
    dong = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    loi = 0

    def ktra(ten: str, ok: bool, them: str = ""):
        nonlocal loi
        print("  %-5s %s" % ("ĐẠT" if ok else "HỎNG", ten))
        if them:
            print("        " + them)
        if not ok:
            loi += 1

    # 1. Đủ bốn phần, đúng thứ tự
    mong = ["PHẦN 1. LÝ THUYẾT VÀ VÍ DỤ MINH HỌA",
            "PHẦN 2. BÀI TẬP TỰ LUYỆN",
            "PHẦN 3. CÂU HỎI TỪ ĐỀ THI",
            "PHẦN 4. HƯỚNG DẪN VÀ ĐÁP ÁN"]
    vi_tri = [next((i for i, t in enumerate(dong) if t == m), -1) for m in mong]
    ktra("Có đủ bốn phần", all(v >= 0 for v in vi_tri),
         "thiếu: " + ", ".join(m for m, v in zip(mong, vi_tri) if v < 0)
         if not all(v >= 0 for v in vi_tri) else "")
    ktra("Bốn phần đúng thứ tự", vi_tri == sorted(vi_tri), str(vi_tri))

    # 2. ĐIỀU QUAN TRỌNG NHẤT: đáp án không lộ trước Phần 4
    if vi_tri[3] > 0:
        truoc = dong[:vi_tri[3]]
        lo = [t for t in truoc if t.startswith(("Đáp án:", "Đáp án ")) or "Đáp án đúng" in t]
        ktra("Không lộ đáp án trước Phần 4", not lo,
             "lộ: " + " | ".join(lo[:3]) if lo else "")

        # Lời giải chỉ được xuất hiện ở ví dụ minh họa của Phần 1
        giai = [i for i, t in enumerate(truoc) if t.startswith("Giải:")]
        trong_phan_1 = all(i < vi_tri[1] for i in giai)
        ktra("Lời giải chỉ nằm ở ví dụ của Phần 1", trong_phan_1,
             "có %d dòng 'Giải:', đều trước Phần 2" % len(giai) if trong_phan_1
             else "có dòng 'Giải:' lọt sang Phần 2 hoặc 3")

    # 3. Dải mức độ ghi theo khoảng câu, không ghi trên từng câu
    dai = [t for t in dong if t.startswith("--- MỨC")]
    ktra("Có dải mức độ theo khoảng câu", len(dai) >= 2, " / ".join(dai))

    # 4. Phần 3 đánh số Câu T kèm thẻ năm
    cau_t = [t for t in dong if t.startswith("Câu T")]
    co_the_nam = all("[" in t and "]" in t for t in cau_t)
    ktra("Phần 3 đánh số 'Câu T' kèm thẻ năm", bool(cau_t) and co_the_nam,
         " | ".join(t[:46] for t in cau_t))

    # 5. Không gán thẻ năm cho câu không rõ xuất xứ.
    # Đếm TRÊN TỆP: chỉ 2 câu trong dữ liệu thử có ghi "Đề thi tốt nghiệp THPT".
    ktra("Chỉ câu truy được nguồn mới vào Phần 3", len(cau_t) == 2,
         "đưa %d câu vào Phần 3, đúng ra là 2" % len(cau_t))

    # 6. Phần 4 dùng bảng ba cột
    bang3 = [b for b in d.tables if len(b.columns) == 3]
    ktra("Phần 4 dùng bảng ba cột", len(bang3) >= 1,
         "có %d bảng, %d bảng ba cột" % (len(d.tables), len(bang3)))
    if bang3:
        dau = [c.text.strip() for c in bang3[0].rows[0].cells]
        ktra("Tiêu đề bảng đúng Câu / Đáp án / Hướng dẫn giải",
             dau == ["Câu", "Đáp án", "Hướng dẫn giải"], str(dau))

    # 7. Mọi câu tự luyện đều có mặt trong bảng đáp án
    if bang3:
        trong_bang = {r.cells[0].text.strip() for b in bang3 for r in b.rows[1:]}
        so_tu_luyen = len(CAU) - 2          # 8 câu, 2 câu có nguồn đề thi
        thieu = [f"Câu {i}" for i in range(1, so_tu_luyen + 1)
                 if f"Câu {i}" not in trong_bang]
        ktra("Bảng đáp án đủ mọi câu tự luyện", not thieu, "thiếu: " + str(thieu[:5]))

    print()
    print("Tệp đã dựng:", RA)
    print("KẾT QUẢ:", "ĐẠT" if loi == 0 else "%d LỖI" % loi)
    return loi


if __name__ == "__main__":
    sys.exit(1 if main() else 0)
