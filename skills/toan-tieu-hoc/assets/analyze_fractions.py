# -*- coding: utf-8 -*-
import re
from docx import Document

DATE_FULL = re.compile(r'\d{1,2}/\d{1,2}/\d{4}')
DATE_DETACHED = re.compile(r'[Nn]gày\s+\d{1,2}/\d{1,2}(?=\s+năm\s+\d{4})')

VNCHAR = r'a-zA-Zàáảãạăằắẳẵặâầấẩẫậèéẻẽẹêềếểễệìíỉĩịòóỏõọôồốổỗộơờớởỡợùúủũụưừứửữựỳýỷỹỵđ' \
         r'ÀÁẢÃẠĂẰẮẲẴẶÂẦẤẨẪẬÈÉẺẼẸÊỀẾỂỄỆÌÍỈĨỊÒÓỎÕỌÔỒỐỔỖỘƠỜỚỞỠỢÙÚỦŨỤƯỪỨỬỮỰỲÝỶỸỴĐ'

# 1) So voi so, co the co ×/+/- va ngoac, nhung KHONG bao gio la chu don le dung mot minh
NUMERIC_TERM = r'\d+(?:×[A-Za-z])?|[A-Za-z]×\d+'
FRACTION_NUMERIC = re.compile(
    r'(?P<num>' + NUMERIC_TERM + r')\s*/\s*'
    r'(?:\((?P<denp>[^()]*?\d[^()]*?)\)|(?P<den>' + NUMERIC_TERM + r'))'
)

# 2) Dang "so/chu" nhu 1/a  -> chi khop khi ben phai la DUY NHAT 1 chu cai, khong dinh lien chu khac
FRACTION_DIGIT_OVER_LETTER = re.compile(
    r'(?P<num>\d+)\s*/\s*(?P<den>[A-Za-z])(?![' + VNCHAR + r'])(?<![' + VNCHAR + r'])'
    .replace("(?<![" + VNCHAR + "])", "")  # placeholder, se sua lookbehind rieng ben duoi
)
FRACTION_DIGIT_OVER_LETTER = re.compile(
    r'(?P<num>\d+)\s*/\s*(?P<den>[A-Za-z])(?![' + VNCHAR + r'])'
)

# 3) Dang "chu/(bieu thuc)" nhu a/(a+8) -> chu don le, KHONG dung lien truoc no la 1 chu cai khac
FRACTION_LETTER_OVER_PAREN = re.compile(
    r'(?<![' + VNCHAR + r'])(?P<num>[A-Za-z])\s*/\s*\((?P<denp>[^()]*?\d[^()]*?)\)'
)

def protected_spans(text):
    spans = []
    for m in DATE_FULL.finditer(text):
        spans.append((m.start(), m.end()))
    for m in DATE_DETACHED.finditer(text):
        spans.append((m.start(), m.end()))
    return spans

def overlaps(pos, spans):
    return any(s <= pos < e for s, e in spans)

def find_all_fractions(text):
    """Tra ve list (start, end, num_text, den_text) khong chong lan, da loai tru ngay thang."""
    prot = protected_spans(text)
    candidates = []
    for regex in (FRACTION_NUMERIC, FRACTION_DIGIT_OVER_LETTER, FRACTION_LETTER_OVER_PAREN):
        for m in regex.finditer(text):
            if overlaps(m.start(), prot):
                continue
            num = m.group('num')
            den = m.groupdict().get('denp') or m.groupdict().get('den')
            candidates.append((m.start(), m.end(), num, den))
    candidates.sort(key=lambda c: (c[0], -(c[1]-c[0])))
    result = []
    last_end = -1
    for s, e, num, den in candidates:
        if s < last_end:
            continue
        result.append((s, e, num, den))
        last_end = e
    return result

if __name__ == "__main__":
    files = [
     (r'D:\Tài liệu số\Toán\Toán 4\5. Học sinh giỏi và Olympic\152 Bậc Thang Chinh Phục Toán 4.docx','B1'),
     (r'D:\Tài liệu số\Toán\Toán 4\5. Học sinh giỏi và Olympic\13 Chuyên Đề Rèn Luyện Tư Duy Tính Toán - Toán 4.docx','B2'),
     (r'D:\Tài liệu số\Toán\Toán 4\5. Học sinh giỏi và Olympic\Phương Pháp Vàng Giải Toán Nâng Cao - Toán 4.docx','B3'),
     (r'D:\Tài liệu số\Toán\Toán 4\5. Học sinh giỏi và Olympic\32 Đề Ôn Luyện Đội Tuyển Toán 4.docx','B4'),
     (r'D:\Tài liệu số\Toán\Toán 4\5. Học sinh giỏi và Olympic\48 Bài Toán Kim Cương Bồi Dưỡng HSG Lớp 4.docx','B5'),
    ]
    total = 0
    uniq = {}
    for path, tag in files:
        d = Document(path)
        texts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    texts.append(c.text)
        for t in texts:
            for s, e, num, den in find_all_fractions(t):
                total += 1
                key = t[s:e]
                uniq.setdefault(key, []).append((tag, t[max(0,s-15):e+10]))
    print("TOTAL:", total, " UNIQUE:", len(uniq))
    print()
    for k in sorted(uniq.keys()):
        tag, ctx = uniq[k][0]
        print(f"{k!r:30s}  [{tag}]  {ctx!r}")
