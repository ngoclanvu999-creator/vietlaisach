# -*- coding: utf-8 -*-
"""
Quet toan bo paragraph/o bang trong 1 file .docx, tim cac phan so dang chu "a/b"
va thay bang cong thuc phan so THAT (OMML - tu/mau xep chong, co gach ngang),
dung chuan hien thi cua sach giao khoa / Word Equation.
Giu nguyen dinh dang (dam/mau chu/size) cua tung run goc.
"""
import copy
import re
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from analyze_fractions import find_all_fractions

MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
def M(tag):
    return f"{{{MATH_NS}}}{tag}"


def _w_rpr_from_run(run):
    """Tra ve mot w:rPr element (hoac None) sao chep dinh dang cua run goc."""
    rPr = run._element.find(qn('w:rPr'))
    if rPr is None:
        return None
    return copy.deepcopy(rPr)


def _make_plain_run(text, rpr_template):
    r_el = etree.Element(qn('w:r'))
    if rpr_template is not None:
        r_el.append(copy.deepcopy(rpr_template))
    t_el = etree.SubElement(r_el, qn('w:t'))
    t_el.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t_el.text = text
    return r_el


def _make_math_text_run(text, rpr_template):
    """m:r chua w:rPr (copy dinh dang) + m:t, tat auto-italic bang m:rPr/m:sty=p."""
    mr = etree.Element(M('r'))
    mrpr = etree.SubElement(mr, M('rPr'))
    sty = etree.SubElement(mrpr, M('sty'))
    sty.set(M('val'), 'p')  # 'p' = plain (khong tu dong in nghieng)
    if rpr_template is not None:
        mr.append(copy.deepcopy(rpr_template))
    mt = etree.SubElement(mr, M('t'))
    mt.text = text
    return mr


def _make_fraction_omath(num_text, den_text, rpr_template):
    omath = etree.Element(M('oMath'))
    f = etree.SubElement(omath, M('f'))
    fpr = etree.SubElement(f, M('fPr'))
    ftype = etree.SubElement(fpr, M('type'))
    ftype.set(M('val'), 'bar')
    ctrlpr = etree.SubElement(fpr, M('ctrlPr'))
    if rpr_template is not None:
        ctrlpr.append(copy.deepcopy(rpr_template))
    num = etree.SubElement(f, M('num'))
    num.append(_make_math_text_run(num_text, rpr_template))
    den = etree.SubElement(f, M('den'))
    den.append(_make_math_text_run(den_text, rpr_template))
    return omath


def rewrite_paragraph_with_fractions(paragraph):
    """Neu paragraph co phan so dang chu, thay the bang OMML. Tra ve True neu co thay doi."""
    p_el = paragraph._p
    runs = paragraph.runs
    if not runs:
        return False
    # Kiem tra truoc: co can xu ly khong?
    any_match = False
    per_run_matches = []
    for run in runs:
        text = run.text
        matches = find_all_fractions(text) if text else []
        per_run_matches.append(matches)
        if matches:
            any_match = True
    if not any_match:
        return False

    # Xoa toan bo w:r hien co khoi paragraph (giu lai w:pPr neu co)
    for run in list(runs):
        p_el.remove(run._element)

    for run, matches in zip(runs, per_run_matches):
        text = run.text
        rpr_template = _w_rpr_from_run(run)
        if not matches:
            if text:
                p_el.append(_make_plain_run(text, rpr_template))
            continue
        pos = 0
        for s, e, num, den in matches:
            if s > pos:
                p_el.append(_make_plain_run(text[pos:s], rpr_template))
            p_el.append(_make_fraction_omath(num, den, rpr_template))
            pos = e
        if pos < len(text):
            p_el.append(_make_plain_run(text[pos:], rpr_template))
    return True


def process_docx(path_in, path_out=None):
    if path_out is None:
        path_out = path_in
    doc = Document(path_in)
    changed = 0
    total_paras = 0

    def process_all_paragraphs(paragraphs):
        nonlocal changed, total_paras
        for p in paragraphs:
            total_paras += 1
            if rewrite_paragraph_with_fractions(p):
                changed += 1

    process_all_paragraphs(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                process_all_paragraphs(cell.paragraphs)

    doc.save(path_out)
    return changed, total_paras


if __name__ == "__main__":
    import sys
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else src
    changed, total = process_docx(src, dst)
    print(f"Paragraphs changed: {changed} / {total}")
