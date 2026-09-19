# -*- coding: utf-8 -*-
"""
Dựng SÁCH biên soạn lại, và rẽ nhánh sang các bộ dựng chuyên biệt.

Thể thức chung (khổ giấy, lề Nghị định 30, kiểu chữ theo cấp học, cách dựng một
câu hỏi) đã tách sang `core/exporter_nen.py` ngày 19/09/2026 — ba ứng dụng đều
dùng chung phần đó. Ở đây chỉ còn bố cục sách và bảng rẽ nhánh.
"""

from pathlib import Path
from typing import Optional, List
import docx
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from core.rewriter import RewrittenBook, RewrittenChapter, RewrittenQuestionItem
from core.loai_dau_ra import la_de as la_loai_de
from core.trang_tri import cach_bay, icon as bt

# Thể thức chung dùng lại từ nền, không định nghĩa lại ở đây
from core.exporter_nen import (
    NenWord, PAPER_SIZES_MM, FONT_MAIN,
    COLOR_PRIMARY, COLOR_SECONDARY, COLOR_SUCCESS, COLOR_WARNING,
    COLOR_TEXT_MAIN, COLOR_TEXT_MUTED,
    set_cell_background, set_callout_borders, add_callout_box,
)


class DocxBookExporter(NenWord):
    @classmethod
    def render_solution_section(cls, doc: docx.Document, book: RewrittenBook, opts: dict):
        """
        Phần lời giải dồn về cuối sách.

        Dùng khi người dùng muốn học sinh tự làm hết bài rồi mới xem đáp án —
        in lời giải ngay dưới đề thì nhìn xuống là thấy, mất tác dụng luyện tập.
        """
        if not opts["solution"] and not opts["casio"] and not opts["traps"] and not opts["answer_key"]:
            return

        doc.add_page_break()
        p_head = doc.add_paragraph()
        p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_head.paragraph_format.space_after = Pt(14)
        r_head = p_head.add_run("PHẦN LỜI GIẢI CHI TIẾT")
        r_head.bold = True
        r_head.font.name = FONT_MAIN
        r_head.font.size = Pt(17)
        r_head.font.color.rgb = COLOR_PRIMARY

        # Gom câu theo chương để lời giải bám đúng bố cục sách
        nhom = []
        if book.chapters:
            for ch in book.chapters:
                nhom.append((ch.title, ch.questions))
        else:
            nhom.append(("", book.questions))

        # Bảng đáp án tra nhanh đặt trước, để dò kết quả mà chưa cần đọc lời giải
        if opts["answer_key"]:
            for tieu_de, ds_cau in nhom:
                co_dap_an = [q for q in ds_cau if (q.correct_answer or "").strip()]
                if not co_dap_an:
                    continue
                if tieu_de:
                    p_ch = doc.add_paragraph()
                    p_ch.paragraph_format.space_before = Pt(10)
                    r_ch = p_ch.add_run(tieu_de)
                    r_ch.bold = True
                    r_ch.font.name = FONT_MAIN
                    r_ch.font.size = Pt(13)
                    r_ch.font.color.rgb = COLOR_SECONDARY
                cls.render_answer_key(doc, ds_cau)

        for tieu_de, ds_cau in nhom:
            if tieu_de:
                p_ch = doc.add_paragraph()
                p_ch.paragraph_format.space_before = Pt(16)
                p_ch.paragraph_format.space_after = Pt(8)
                r_ch = p_ch.add_run(tieu_de)
                r_ch.bold = True
                r_ch.font.name = FONT_MAIN
                r_ch.font.size = Pt(14)
                r_ch.font.color.rgb = COLOR_PRIMARY

            for q in ds_cau:
                p_t = doc.add_paragraph()
                p_t.paragraph_format.space_before = Pt(10)
                p_t.paragraph_format.space_after = Pt(3)
                r_t = p_t.add_run(f"{q.title}. ")
                r_t.bold = True
                r_t.font.name = FONT_MAIN
                r_t.font.size = Pt(13)
                r_t.font.color.rgb = COLOR_PRIMARY
                if q.correct_answer:
                    r_da = p_t.add_run(f"Đáp án {q.correct_answer.strip().upper()[:1]}")
                    r_da.bold = True
                    r_da.font.name = FONT_MAIN
                    r_da.font.size = Pt(12.5)
                    r_da.font.color.rgb = COLOR_SUCCESS

                if q.solution_method1 and opts["solution"]:
                    for line in q.solution_method1.strip().split("\n"):
                        if not line.strip():
                            continue
                        p_l = doc.add_paragraph()
                        p_l.paragraph_format.space_after = Pt(2)
                        p_l.paragraph_format.line_spacing = 1.25
                        r_l = p_l.add_run(line.strip())
                        r_l.font.name = FONT_MAIN
                        r_l.font.size = Pt(12.5)
                        r_l.font.color.rgb = COLOR_TEXT_MAIN

                if q.solution_method2 and opts["casio"]:
                    add_callout_box(doc, "💡 KỸ THUẬT BẤM MÁY CASIO FX-580VN X",
                                    q.solution_method2, "F0FDF4", "2F855A", COLOR_SUCCESS)
                if q.trap_warning and opts["traps"]:
                    add_callout_box(doc, "⚠️ BẪY ĐỀ THI & LỖI SAI HỌC SINH HAY MẮC",
                                    q.trap_warning, "FFF5F5", "C53030", COLOR_WARNING)


    # =======================================================================
    # KIỂU RIÊNG CHO ĐỀ THI
    # Đề thi phải đọc được như một đề thi thật: học sinh làm bài trước, xem đáp
    # án sau. Chen lời giải và mẹo Casio ngay dưới từng câu là hỏng mục đích sử
    # dụng — nhìn xuống là thấy đáp án, không còn gì để luyện.
    # =======================================================================

    @classmethod
    def render_exam_header(cls, doc: docx.Document, book: RewrittenBook):
        """Dựng lại phần đầu đề thi theo thể thức hành chính quen thuộc."""
        info = book.exam_info or {}

        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width = Inches(3.2)
        table.columns[1].width = Inches(3.3)

        trai = table.cell(0, 0)
        p_trai = trai.paragraphs[0]
        p_trai.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_trai.add_run(info.get("don_vi") or "ĐƠN VỊ TỔ CHỨC")
        r.bold = True
        r.font.name = FONT_MAIN
        r.font.size = Pt(12)
        p2 = trai.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run("(Đề thi gồm nhiều trang)")
        r2.italic = True
        r2.font.name = FONT_MAIN
        r2.font.size = Pt(10)
        r2.font.color.rgb = COLOR_TEXT_MUTED

        phai = table.cell(0, 1)
        p_phai = phai.paragraphs[0]
        p_phai.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p_phai.add_run(info.get("ky_thi") or book.new_title.upper())
        r3.bold = True
        r3.font.name = FONT_MAIN
        r3.font.size = Pt(13)
        r3.font.color.rgb = COLOR_PRIMARY

        dong_phu = []
        if info.get("mon"):
            dong_phu.append(f"Môn: {info['mon']}")
        if info.get("thoi_gian"):
            dong_phu.append(f"Thời gian làm bài: {info['thoi_gian']}")
        if not dong_phu:
            dong_phu.append("Thời gian làm bài: 90 phút (không kể thời gian phát đề)")
        p4 = phai.add_paragraph()
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r4 = p4.add_run(" — ".join(dong_phu))
        r4.italic = True
        r4.font.name = FONT_MAIN
        r4.font.size = Pt(11)

        if info.get("ma_de"):
            p5 = phai.add_paragraph()
            p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r5 = p5.add_run(f"Mã đề: {info['ma_de']}")
            r5.bold = True
            r5.font.name = FONT_MAIN
            r5.font.size = Pt(11.5)

        # Dòng họ tên và số báo danh
        p_hs = doc.add_paragraph()
        p_hs.paragraph_format.space_before = Pt(14)
        p_hs.paragraph_format.space_after = Pt(12)
        r_hs = p_hs.add_run("Họ và tên thí sinh: " + "." * 42 + "   Số báo danh: " + "." * 14)
        r_hs.font.name = FONT_MAIN
        r_hs.font.size = Pt(12.5)

    @classmethod
    def render_exam_question(cls, doc: docx.Document, q: RewrittenQuestionItem, so_thu_tu: int):
        """Câu hỏi trong phần đề: chỉ đề bài và 4 phương án, KHÔNG lộ đáp án."""
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(8)
        p_q.paragraph_format.space_after = Pt(3)
        p_q.paragraph_format.line_spacing = 1.25

        r_num = p_q.add_run(f"Câu {so_thu_tu}. ")
        r_num.bold = True
        r_num.font.name = FONT_MAIN
        r_num.font.size = Pt(13)
        r_num.font.color.rgb = COLOR_PRIMARY

        r_ct = p_q.add_run(q.new_content)
        r_ct.font.name = FONT_MAIN
        r_ct.font.size = Pt(13)
        r_ct.font.color.rgb = COLOR_TEXT_MAIN

        if not q.new_options:
            return

        # Bốn phương án dàn 2x2, KHÔNG tô đậm đáp án đúng
        table = doc.add_table(rows=2, cols=2)
        table.autofit = False
        for row in table.rows:
            for cell in row.cells:
                cell.width = Inches(3.25)

        for i, opt in enumerate(q.new_options[:4]):
            c = table.cell(i // 2, i % 2)
            p_opt = c.paragraphs[0]
            p_opt.paragraph_format.space_before = Pt(1)
            p_opt.paragraph_format.space_after = Pt(1)
            r_opt = p_opt.add_run(opt)
            r_opt.font.name = FONT_MAIN
            r_opt.font.size = Pt(12.5)
            r_opt.font.color.rgb = COLOR_TEXT_MAIN

    @classmethod
    def render_answer_key(cls, doc: docx.Document, questions: list):
        """Bảng đáp án 10 câu mỗi hàng, tra cứu nhanh như đáp án đề thi thật."""
        p_head = doc.add_paragraph()
        p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_head.paragraph_format.space_after = Pt(12)
        r = p_head.add_run("BẢNG ĐÁP ÁN")
        r.bold = True
        r.font.name = FONT_MAIN
        r.font.size = Pt(16)
        r.font.color.rgb = COLOR_PRIMARY

        moi_hang = 10
        for bat_dau in range(0, len(questions), moi_hang):
            nhom = questions[bat_dau:bat_dau + moi_hang]
            table = doc.add_table(rows=2, cols=len(nhom))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            for i, q in enumerate(nhom):
                o_cau = table.cell(0, i)
                set_cell_background(o_cau, "EFF6FF")
                p_c = o_cau.paragraphs[0]
                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_c = p_c.add_run(str(bat_dau + i + 1))
                r_c.bold = True
                r_c.font.name = FONT_MAIN
                r_c.font.size = Pt(11)
                r_c.font.color.rgb = COLOR_PRIMARY

                o_da = table.cell(1, i)
                p_d = o_da.paragraphs[0]
                p_d.alignment = WD_ALIGN_PARAGRAPH.CENTER
                dap_an = (q.correct_answer or "—").strip().upper()[:1] or "—"
                r_d = p_d.add_run(dap_an)
                r_d.bold = True
                r_d.font.name = FONT_MAIN
                r_d.font.size = Pt(12)
                r_d.font.color.rgb = COLOR_SUCCESS

            doc.add_paragraph().paragraph_format.space_after = Pt(6)

    @classmethod
    def export_exam(cls, book: RewrittenBook, doc: docx.Document, opts: dict):
        """Bố cục đề thi: phần đề → bảng đáp án → hướng dẫn giải chi tiết."""
        ds_cau = book.questions or [q for ch in book.chapters for q in ch.questions]

        cls.render_exam_header(doc, book)

        for i, q in enumerate(ds_cau, 1):
            cls.render_exam_question(doc, q, i)

        p_het = doc.add_paragraph()
        p_het.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_het.paragraph_format.space_before = Pt(16)
        r_het = p_het.add_run("---------- HẾT ----------")
        r_het.bold = True
        r_het.font.name = FONT_MAIN
        r_het.font.size = Pt(12)

        p_note = doc.add_paragraph()
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_note = p_note.add_run("Thí sinh không được sử dụng tài liệu. Cán bộ coi thi không giải thích gì thêm.")
        r_note.italic = True
        r_note.font.name = FONT_MAIN
        r_note.font.size = Pt(10.5)
        r_note.font.color.rgb = COLOR_TEXT_MUTED

        # ---- Bảng đáp án ----
        doc.add_page_break()
        cls.render_answer_key(doc, ds_cau)

        # ---- Hướng dẫn giải chi tiết ----
        doc.add_page_break()
        p_hd = doc.add_paragraph()
        p_hd.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_hd.paragraph_format.space_after = Pt(14)
        r_hd = p_hd.add_run("HƯỚNG DẪN GIẢI CHI TIẾT")
        r_hd.bold = True
        r_hd.font.name = FONT_MAIN
        r_hd.font.size = Pt(16)
        r_hd.font.color.rgb = COLOR_PRIMARY

        for i, q in enumerate(ds_cau, 1):
            p_t = doc.add_paragraph()
            p_t.paragraph_format.space_before = Pt(10)
            p_t.paragraph_format.space_after = Pt(3)
            r_t = p_t.add_run(f"Câu {i}. ")
            r_t.bold = True
            r_t.font.name = FONT_MAIN
            r_t.font.size = Pt(13)
            r_t.font.color.rgb = COLOR_PRIMARY
            if q.correct_answer:
                r_da = p_t.add_run(f"Đáp án {q.correct_answer.strip().upper()[:1]}")
                r_da.bold = True
                r_da.font.name = FONT_MAIN
                r_da.font.size = Pt(12.5)
                r_da.font.color.rgb = COLOR_SUCCESS

            if q.solution_method1:
                for line in q.solution_method1.strip().split("\n"):
                    if not line.strip():
                        continue
                    p_l = doc.add_paragraph()
                    p_l.paragraph_format.space_after = Pt(2)
                    p_l.paragraph_format.line_spacing = 1.2
                    r_l = p_l.add_run(line.strip())
                    r_l.font.name = FONT_MAIN
                    r_l.font.size = Pt(12.5)
                    r_l.font.color.rgb = COLOR_TEXT_MAIN

            if q.solution_method2 and opts["casio"]:
                add_callout_box(doc, "💡 KỸ THUẬT CASIO fx-580VN X",
                                q.solution_method2, "F0FDF4", "2F855A", COLOR_SUCCESS)
            if q.trap_warning and opts["traps"]:
                add_callout_box(doc, "⚠️ BẪY & LỖI SAI THƯỜNG GẶP",
                                q.trap_warning, "FFF5F5", "C53030", COLOR_WARNING)

        return doc

    @classmethod
    def export(
        cls,
        book: RewrittenBook,
        output_path: Path,
        paper_format: str = "a4",
        options: Optional[dict] = None
    ) -> Path:
        opts = cls.merge_options(options)
        loai_ra_som = getattr(book, "loai_dau_ra", "") or ""
        cb = cach_bay(getattr(book, "cap_hoc", "THPT"), loai_ra_som)

        doc = docx.Document()
        cls.apply_default_style(doc, cb)
        section = doc.sections[0]
        cls.apply_standard_page_setup(section, paper_format=paper_format)

        # Header & Footer chuẩn công văn
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run(f"TÀI LIỆU LƯU HÀNH NỘI BỘ - CHƯƠNG TRÌNH GDPT 2018 | {book.new_title}")
        r_hdr.font.name = FONT_MAIN
        r_hdr.font.size = Pt(9)
        r_hdr.font.color.rgb = COLOR_TEXT_MUTED

        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ftr = p_ftr.add_run("Trang được biên soạn & tối ưu hóa tự động - Chuẩn thể thức Bộ GD&ĐT")
        r_ftr.font.name = FONT_MAIN
        r_ftr.font.size = Pt(9)
        r_ftr.font.color.rgb = COLOR_TEXT_MUTED

        # ==========================================
        # RẼ NHÁNH THEO LOẠI TÀI LIỆU
        # Đầu vào là đề thi thì đầu ra cũng phải là đề thi: học sinh làm bài
        # trước rồi mới xem đáp án. Đầu vào là sách thì giữ bố cục sách.
        # ==========================================
        # Chín loại đầu ra: năm loại đề đi qua bộ dựng chuẩn 2025 (ma trận, bản
        # đặc tả, đề ba phần, hướng dẫn chấm). Các loại còn lại giữ bố cục cũ.
        loai_ra = getattr(book, "loai_dau_ra", "") or ""
        mon = getattr(book, "subject", "") or "toan"
        if la_loai_de(loai_ra):
            from core.exporter_de import xuat_bo_de
            xuat_bo_de(doc, book, loai_ra, mon,
                       kem_loi_giai=bool(opts.get("solution", True)))
            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(output_path))
            return output_path

        # Tài liệu HSG không dùng chung bộ dựng sách: đề HSG là tự luận, phải
        # giấu phương án đi chứ không cho sẵn bốn lựa chọn để đoán.
        if loai_ra == "TAI_LIEU_HSG":
            from core.exporter_hsg import xuat_tai_lieu_hsg
            xuat_tai_lieu_hsg(doc, book, mon, opts)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(output_path))
            return output_path

        # Chuyên đề bài tập: khuôn BỐN PHẦN, đáp án dồn hết về Phần 4. Trước đây
        # loại này không có nhánh riêng nên rơi xuống bộ dựng sách bên dưới, mà
        # bộ đó in lời giải ngay dưới mỗi bài — trái hẳn skill chuyen-de-bai-tap.
        if loai_ra == "CHUYEN_DE_BT":
            from core.exporter_chuyen_de import xuat_chuyen_de
            xuat_chuyen_de(doc, book, mon, opts)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(output_path))
            return output_path

        # Giao an: khung Cong van 5512. La ho so chuyen mon nen o MOI cap hoc
        # deu giu nghiem ngat, khong bieu tuong, khong trang tri.
        if loai_ra == "GIAO_AN":
            from core.exporter_giao_an import xuat_giao_an
            xuat_giao_an(doc, book, mon,
                         cap_hoc=getattr(book, "cap_hoc", "THPT"),
                         ten_bai=getattr(book, "ten_bai", "") or "",
                         options=opts)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(output_path))
            return output_path

        if getattr(book, "doc_type", "") == "DE_THI":
            cls.export_exam(book, doc, opts)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(output_path))
            return output_path

        # ==========================================
        # TRANG TIÊU ĐỀ & LỜI TỰA
        # ==========================================
        # Tiểu học: dòng điền tên ngay đầu tài liệu, theo đúng thói quen của
        # phiếu bài tập — trẻ nhận phiếu là điền tên trước khi làm.
        if cb.co_dong_ho_ten:
            p_ht = doc.add_paragraph()
            p_ht.paragraph_format.space_after = Pt(10)
            r_ht = p_ht.add_run("Họ và tên: ………………………………………     Lớp: …………")
            r_ht.bold = True
            r_ht.font.name = FONT_MAIN
            r_ht.font.size = Pt(cb.co_chu)
            r_ht.font.color.rgb = COLOR_TEXT_MAIN

        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(30)

        # Tên sách chuẩn hóa
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(12)
        r_title = p_title.add_run(book.new_title.upper())
        r_title.bold = True
        r_title.font.name = FONT_MAIN
        r_title.font.size = Pt(20)
        r_title.font.color.rgb = COLOR_PRIMARY

        # Phụ đề
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(20)
        r_sub = p_sub.add_run(book.subtitle)
        r_sub.italic = True
        r_sub.font.name = FONT_MAIN
        r_sub.font.size = Pt(13)
        r_sub.font.color.rgb = COLOR_SECONDARY

        # Lời tựa truyền cảm hứng
        if book.author_note and opts["foreword"]:
            add_callout_box(
                doc,
                title="✨ LỜI TỰA TRUYỀN CẢM HỨNG & SỨ MỆNH CUỐN SÁCH",
                content=book.author_note,
                bg_hex="F8FAFC",
                border_hex="2B6CB0",
                title_color=COLOR_PRIMARY
            )

        # Bí kíp thủ khoa
        if opts["secrets"] and getattr(book, "valedictorian_secrets", None):
            secrets_content = "\n".join(f"• {sec}" for sec in book.valedictorian_secrets)
            add_callout_box(
                doc,
                title="🏆 BÍ KÍP VÀNG PHÒNG THI & CHIẾN THUẬT TỪ THỦ KHOA",
                content=secrets_content,
                bg_hex="FFFDF5",
                border_hex="D69E2E",
                title_color=COLOR_WARNING
            )

        # Góc kết nối STEM thực tế đời sống
        if opts["stem"] and getattr(book, "stem_connection", ""):
            add_callout_box(
                doc,
                title="🌐 GÓC KẾT NỐI THỰC TIỄN ĐỜI SỐNG & CÔNG NGHỆ (GDPT 2018)",
                content=book.stem_connection,
                bg_hex="F0FFF4",
                border_hex="38A169",
                title_color=COLOR_SUCCESS
            )


        # ==========================================
        # KIỂU 1: ĐẠI CẨM NANG GỘP TỪ NHIỀU CHƯƠNG (MASTER BOOK)
        # ==========================================
        if book.chapters:
            # Mục lục tổng quan
            p_toc_head = doc.add_paragraph()
            p_toc_head.paragraph_format.space_before = Pt(14)
            p_toc_head.paragraph_format.space_after = Pt(10)
            r_th = p_toc_head.add_run("MỤC LỤC CÁC CHUYÊN ĐỀ")
            r_th.bold = True
            r_th.font.name = FONT_MAIN
            r_th.font.size = Pt(15)
            r_th.font.color.rgb = COLOR_PRIMARY

            for ch in book.chapters:
                p_ch_item = doc.add_paragraph()
                p_ch_item.paragraph_format.space_after = Pt(4)
                r_ci = p_ch_item.add_run(f"• {ch.title} ({len(ch.questions)} bài toán trọng tâm)")
                r_ci.font.name = FONT_MAIN
                r_ci.font.size = Pt(12)

            doc.add_page_break()

            # Lặp qua từng chương
            for ch in book.chapters:
                # Tiêu đề chương
                p_ch_title = doc.add_paragraph()
                p_ch_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_ch_title.paragraph_format.space_before = Pt(18)
                p_ch_title.paragraph_format.space_after = Pt(14)
                r_ct = p_ch_title.add_run(ch.title)
                r_ct.bold = True
                r_ct.font.name = FONT_MAIN
                r_ct.font.size = Pt(16)
                r_ct.font.color.rgb = COLOR_PRIMARY

                # PHẦN I: LÝ THUYẾT TRỌNG TÂM CỦA CHƯƠNG
                if ch.theory_section and opts["theory"]:
                    add_callout_box(
                        doc,
                        title="PHẦN I: KIẾN THỨC TRỌNG TÂM & LÝ THUYẾT CỐT LÕI",
                        content=ch.theory_section,
                        bg_hex="EFF6FF",
                        border_hex="1A365D",
                        title_color=COLOR_PRIMARY
                    )

                # PHẦN II: HỆ THỐNG BÀI TOÁN RÈN LUYỆN
                p_sec_title = doc.add_paragraph()
                p_sec_title.paragraph_format.space_before = Pt(12)
                p_sec_title.paragraph_format.space_after = Pt(10)
                r_st = p_sec_title.add_run("PHẦN II: HỆ THỐNG BÀI TOÁN CHỌN LỌC & LỜI GIẢI ĐA CHIỀU")
                r_st.bold = True
                r_st.font.name = FONT_MAIN
                r_st.font.size = Pt(14)
                r_st.font.color.rgb = COLOR_PRIMARY

                for q in ch.questions:
                    cls.render_question_item(doc, q, options=opts, cb=cb)

                doc.add_page_break()

        # ==========================================
        # KIỂU 2: CUỐN SÁCH ĐƠN (SINGLE BOOK)
        # ==========================================
        else:
            # PHẦN I: KIẾN THỨC TRỌNG TÂM & LÝ THUYẾT CỐT LÕI
            if book.theory_section and opts["theory"]:
                add_callout_box(
                    doc,
                    title="PHẦN I: KIẾN THỨC TRỌNG TÂM & LÝ THUYẾT NỀN TẢNG CHUẨN BGD",
                    content=book.theory_section,
                    bg_hex="EFF6FF",
                    border_hex="1A365D",
                    title_color=COLOR_PRIMARY
                )

            doc.add_page_break()

            # PHẦN II: HỆ THỐNG BÀI TOÁN VÀ LỜI GIẢI ĐA CHIỀU
            p_sec_title = doc.add_paragraph()
            p_sec_title.paragraph_format.space_before = Pt(10)
            p_sec_title.paragraph_format.space_after = Pt(14)
            r_st = p_sec_title.add_run("PHẦN II: HỆ THỐNG BÀI TOÁN RÈN LUYỆN THEO CẤP ĐỘ NHẬN THỨC")
            r_st.bold = True
            r_st.font.name = FONT_MAIN
            r_st.font.size = Pt(14)
            r_st.font.color.rgb = COLOR_PRIMARY

            for q in book.questions:
                cls.render_question_item(doc, q, options=opts, cb=cb)

        if opts["vi_tri_loi_giai"] == cls.LOI_GIAI_CUOI_SACH:
            cls.render_solution_section(doc, book, opts)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path
