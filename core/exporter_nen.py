# -*- coding: utf-8 -*-
"""
Nền dựng tài liệu Word: thể thức chung mà MỌI loại đầu ra đều phải theo.

Tách khỏi `core/exporter.py` ngày 19/09/2026, khi chủ dự án chốt chia công cụ
thành ba ứng dụng riêng. Tệp cũ trộn ba thứ vào một chỗ: thể thức chung, bố cục
sách, và một bảng rẽ nhánh theo loại đầu ra. Bảng rẽ nhánh chính là thứ làm mọi
thứ rối — sửa một loại là phải mở tệp mà bốn loại kia cùng nằm trong đó.

Ở đây CHỈ giữ phần không phụ thuộc loại đầu ra:

  - Khổ giấy và lề theo Nghị định 30/2020/NĐ-CP (lề trái 30mm để đóng gáy)
  - Kiểu chữ mặc định theo cấp học
  - Dựng một câu hỏi kèm phương án, lời giải, cảnh báo bẫy
  - Hộp ghi chú và đường phân cách

Ba app kế thừa lớp `NenWord` nên mọi lời gọi `cls.xxx` cũ vẫn chạy nguyên.
"""

from pathlib import Path
from typing import Optional, List
import docx
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from core.rewriter import RewrittenBook, RewrittenChapter, RewrittenQuestionItem
from core.trang_tri import cach_bay, icon as bt

# Bảng màu sắc chuẩn sư phạm & thanh lịch
COLOR_PRIMARY = RGBColor(26, 54, 93)       # Xanh Navy Đậm #1A365D (Chuẩn sách Bộ GD)
COLOR_SECONDARY = RGBColor(43, 108, 176)   # Xanh Dương #2B6CB0
COLOR_SUCCESS = RGBColor(47, 133, 90)      # Xanh Lá Sư Phạm #2F855A
COLOR_WARNING = RGBColor(197, 48, 48)      # Đỏ Cảnh Báo Bẫy #C53030
COLOR_TEXT_MAIN = RGBColor(26, 32, 44)     # Đen Than #1A202C (Dễ đọc, không chói)
COLOR_TEXT_MUTED = RGBColor(74, 85, 104)   # Xám #4A5568

FONT_MAIN = "Times New Roman"

def set_cell_background(cell, fill_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_callout_borders(cell, border_color_hex: str, border_size_pt: int = 18):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="single" w:sz="{border_size_pt}" w:space="0" w:color="{border_color_hex}"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def add_callout_box(doc: docx.Document, title: str, content: str, bg_hex: str, border_hex: str, title_color: RGBColor):
    """Chèn hộp ghi chú sư phạm chuẩn mực"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)

    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_callout_borders(cell, border_hex, border_size_pt=18)

    p_title = cell.paragraphs[0]
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run(title)
    run_title.bold = True
    run_title.font.name = FONT_MAIN
    run_title.font.size = Pt(11.5)
    run_title.font.color.rgb = title_color

    for line in content.strip().split("\n"):
        if not line.strip():
            continue
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.2
        run = p.add_run(line.strip())
        run.font.name = FONT_MAIN
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_TEXT_MAIN

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)


# Kích thước khổ giấy được hỗ trợ (rộng x cao, đơn vị milimét)
PAPER_SIZES_MM = {
    "a4": (210, 297),
    "b5": (176, 250),
}


class NenWord:
    """Thể thức Word dùng chung. Ba app đều kế thừa lớp này."""

    @classmethod
    def apply_standard_page_setup(cls, section, paper_format: str = "a4"):
        """
        Thiết lập thể thức văn bản chuẩn Nghị định 30/2020/NĐ-CP.
        Mặc định khổ A4 với lề Trái 30mm (chừa gáy sách), Phải 15mm, Trên/Dưới 20mm.
        Khổ B5 dùng khi in sách bỏ túi.
        """
        width_mm, height_mm = PAPER_SIZES_MM.get((paper_format or "a4").lower(), PAPER_SIZES_MM["a4"])
        section.page_width = Mm(width_mm)
        section.page_height = Mm(height_mm)

        # Lề trang: Trái 30mm (đóng gáy sách), Phải 15mm, Trên 20mm, Dưới 20mm
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)

    @classmethod
    def apply_default_style(cls, doc, cb=None):
        """
        Đặt font mặc định cho toàn tài liệu. Nếu không làm bước này, mọi đoạn văn
        không được gán font tường minh sẽ rơi về Calibri 11pt của Word, sai thể thức.

        Cỡ chữ và giãn dòng đổi theo cấp học: tài liệu tiểu học phải to và
        thoáng hơn, trẻ mới đọc được.
        """
        if cb is None:
            cb = cach_bay()
        try:
            style = doc.styles["Normal"]
            style.font.name = FONT_MAIN
            style.font.size = Pt(cb.co_chu)
            # Word dùng thuộc tính riêng cho bảng mã Đông Á, phải đặt kèm mới ăn chắc
            rpr = style.element.get_or_add_rPr()
            rfonts = rpr.get_or_add_rFonts()
            rfonts.set(qn("w:eastAsia"), FONT_MAIN)
            rfonts.set(qn("w:cs"), FONT_MAIN)
            style.paragraph_format.line_spacing = cb.gian_dong
        except Exception:
            pass

    # Mặc định bật hết. Người dùng tắt mục nào thì mục đó biến mất khỏi file Word.
    # Vị trí phần lời giải trong sách
    LOI_GIAI_SAU_MOI_BAI = "sau_moi_bai"   # in ngay dưới từng câu
    LOI_GIAI_CUOI_SACH = "cuoi_sach"       # dồn hết về cuối, học sinh tự làm trước

    DEFAULT_OPTIONS = {
        "theory": True,      # Phần I: lý thuyết nền tảng
        "foreword": True,    # Lời tựa truyền cảm hứng
        "secrets": True,     # Bí kíp thủ khoa
        "stem": True,        # Góc kết nối STEM
        "solution": True,    # Lời giải tự luận chi tiết
        "casio": True,       # Mẹo Casio
        "traps": True,       # Cảnh báo bẫy
        "answer_key": True,  # Bảng đáp án tra nhanh (khi lời giải dồn về cuối)
    }

    @classmethod
    def merge_options(cls, options: Optional[dict]) -> dict:
        merged = dict(cls.DEFAULT_OPTIONS)
        vi_tri = cls.LOI_GIAI_SAU_MOI_BAI
        if options:
            merged.update({k: bool(v) for k, v in options.items() if k in merged})
            gt = str(options.get("vi_tri_loi_giai", "") or "").strip()
            if gt in (cls.LOI_GIAI_SAU_MOI_BAI, cls.LOI_GIAI_CUOI_SACH):
                vi_tri = gt
        merged["vi_tri_loi_giai"] = vi_tri
        return merged

    @classmethod
    def render_question_item(cls, doc: docx.Document, q: RewrittenQuestionItem,
                             options: Optional[dict] = None, cb=None):
        """
        Trình bày từng câu hỏi theo chuẩn sư phạm Bộ GD&ĐT.

        `cb` là bộ tham số trình bày theo cấp học (core/trang_tri.py). Tiểu học
        chữ to hơn, thoáng hơn, có biểu tượng và chừa dòng kẻ để viết bài.
        """
        opts = cls.merge_options(options)
        if cb is None:
            cb = cach_bay()

        # Đề mục câu hỏi
        p_qtitle = doc.add_paragraph()
        p_qtitle.paragraph_format.space_before = Pt(cb.cach_bai)
        p_qtitle.paragraph_format.space_after = Pt(3)

        dau = f"{cb.dau_bai} " if cb.dau_bai else ""
        r_num = p_qtitle.add_run(f"{dau}{q.title}")
        r_num.bold = True
        r_num.font.name = FONT_MAIN
        r_num.font.size = Pt(cb.co_chu)
        r_num.font.color.rgb = cb.mau if not q.is_added_new else COLOR_SUCCESS

        if q.is_added_new:
            r_badge = p_qtitle.add_run(" [MỚI BỔ SUNG]")
            r_badge.bold = True
            r_badge.font.size = Pt(10)
            r_badge.font.color.rgb = COLOR_SUCCESS

        # Nội dung đề bài
        p_qcontent = doc.add_paragraph()
        p_qcontent.paragraph_format.space_before = Pt(2)
        p_qcontent.paragraph_format.space_after = Pt(5)
        p_qcontent.paragraph_format.line_spacing = cb.gian_dong
        r_qc = p_qcontent.add_run(q.new_content)
        r_qc.font.name = FONT_MAIN
        r_qc.font.size = Pt(cb.co_de_bai)
        r_qc.font.color.rgb = COLOR_TEXT_MAIN

        # Các phương án trắc nghiệm A, B, C, D (dàn ô đều)
        if q.new_options:
            table_opt = doc.add_table(rows=2, cols=2)
            table_opt.alignment = WD_TABLE_ALIGNMENT.CENTER
            table_opt.autofit = False
            col_width = Inches(3.1)
            for row in table_opt.rows:
                for cell in row.cells:
                    cell.width = col_width

            for o_idx, opt_text in enumerate(q.new_options[:4]):
                row_i = o_idx // 2
                col_i = o_idx % 2
                c = table_opt.cell(row_i, col_i)
                set_cell_background(c, "F7FAFC")
                p_opt = c.paragraphs[0]
                p_opt.paragraph_format.space_before = Pt(3)
                p_opt.paragraph_format.space_after = Pt(3)
                r_opt = p_opt.add_run(opt_text)
                r_opt.font.name = FONT_MAIN
                r_opt.font.size = Pt(12)
                if q.correct_answer and opt_text.strip().startswith(q.correct_answer):
                    r_opt.bold = True
                    r_opt.font.color.rgb = COLOR_SUCCESS

            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.space_before = Pt(0)
            p_sp.paragraph_format.space_after = Pt(4)

        # Ba phần dưới đây chỉ in kèm câu hỏi khi người dùng chọn để lời giải
        # NGAY SAU MỖI BÀI. Nếu chọn dồn về cuối sách thì bỏ qua hết ở đây,
        # phần render_solution_section() sẽ lo.
        if opts["vi_tri_loi_giai"] == cls.LOI_GIAI_CUOI_SACH:
            cls.render_separator(doc)
            return

        # Lời giải Tự luận chuẩn mực sư phạm
        if q.solution_method1 and opts["solution"]:
            p_sol1_head = doc.add_paragraph()
            p_sol1_head.paragraph_format.space_before = Pt(6)
            p_sol1_head.paragraph_format.space_after = Pt(2)
            r_s1_h = p_sol1_head.add_run("✎ Hướng dẫn giải chi tiết (Tự luận chuẩn mực):")
            r_s1_h.bold = True
            r_s1_h.italic = True
            r_s1_h.font.name = FONT_MAIN
            r_s1_h.font.size = Pt(12.5)
            r_s1_h.font.color.rgb = COLOR_SECONDARY

            for line in q.solution_method1.strip().split("\n"):
                if not line.strip():
                    continue
                p_sol_l = doc.add_paragraph()
                p_sol_l.paragraph_format.space_before = Pt(1)
                p_sol_l.paragraph_format.space_after = Pt(3)
                p_sol_l.paragraph_format.line_spacing = 1.25
                r_sol_l = p_sol_l.add_run(line.strip())
                r_sol_l.font.name = FONT_MAIN
                r_sol_l.font.size = Pt(12.5)
                r_sol_l.font.color.rgb = COLOR_TEXT_MAIN

        # Khung Kỹ thuật Casio & Mẹo nhanh
        if q.solution_method2 and opts["casio"]:
            add_callout_box(
                doc,
                title="💡 KỸ THUẬT BẤM MÁY CASIO FX-580VN X & TƯ DUY GIẢI NHANH",
                content=q.solution_method2,
                bg_hex="F0FDF4",
                border_hex="2F855A",
                title_color=COLOR_SUCCESS
            )

        # Khung Cảnh báo bẫy sai lầm
        if q.trap_warning and opts["traps"]:
            add_callout_box(
                doc,
                title="⚠️ BẪY ĐỀ THI & LỖI SAI HỌC SINH HAY MẮC PHẢI",
                content=q.trap_warning,
                bg_hex="FFF5F5",
                border_hex="C53030",
                title_color=COLOR_WARNING
            )

        # Tiểu học: chừa dòng kẻ để trẻ viết bài giải ngay vào tài liệu in ra.
        # Đây là thói quen của phiếu bài tập tiểu học, không dùng ở THPT.
        if cb.dong_ke_lam_bai and not opts.get("solution", True):
            p_bg = doc.add_paragraph()
            p_bg.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_bg.paragraph_format.space_before = Pt(4)
            p_bg.paragraph_format.space_after = Pt(2)
            r_bg = p_bg.add_run("BÀI GIẢI")
            r_bg.bold = True
            r_bg.font.name = FONT_MAIN
            r_bg.font.size = Pt(cb.co_chu - 1)
            r_bg.font.color.rgb = cb.mau
            for _ in range(cb.dong_ke_lam_bai):
                p_ke = doc.add_paragraph()
                p_ke.paragraph_format.space_before = Pt(0)
                p_ke.paragraph_format.space_after = Pt(6)
                r_ke = p_ke.add_run("." * 78)
                r_ke.font.name = FONT_MAIN
                r_ke.font.size = Pt(cb.co_chu)
                r_ke.font.color.rgb = COLOR_TEXT_MUTED

        cls.render_separator(doc)

    @classmethod
    def render_separator(cls, doc: docx.Document):
        """Đường kẻ phân cách nhẹ giữa các bài."""
        p_sep = doc.add_paragraph()
        p_sep.paragraph_format.space_before = Pt(4)
        p_sep.paragraph_format.space_after = Pt(8)
        p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_sep = p_sep.add_run("· · · — — — · · ·")
        r_sep.font.color.rgb = RGBColor(203, 213, 225)
        r_sep.font.size = Pt(8)

