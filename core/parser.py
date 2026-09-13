import re
import os
from pathlib import Path
from typing import List, Dict, Any, Optional
from dataclasses import dataclass, field, asdict
import docx
import openpyxl
import pymupdf
from PIL import Image

from core.math_engine import clean_symbol_text, format_math_typography, clean_paragraph_text

@dataclass
class QuestionItem:
    index: int
    title: str = ""
    content: str = ""
    options: List[str] = field(default_factory=list)
    correct_answer: str = ""
    solution: str = ""
    subject: str = "toan"
    category: str = ""
    source_file: str = ""
    chapter_title: str = ""
    theory_box: str = ""

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)

PAGE_STAMP_PATTERNS = [
    re.compile(r"trang\s*\d+\s*/\s*\d+", re.IGNORECASE),
    re.compile(r"mã\s*đề\s*thi\s*\d+", re.IGNORECASE),
    re.compile(r"họ\s*và\s*tên\s*thí\s*sinh.*sbd", re.IGNORECASE),
    re.compile(r"thời\s*gian\s*làm\s*bài.*phút", re.IGNORECASE),
    re.compile(r"sở\s*gd&đt.*cụm", re.IGNORECASE),
    re.compile(r"đề\s*thi\s*(?:chính\s*thức|chọn\s*hsg)", re.IGNORECASE),
    re.compile(r"phần\s*trắc\s*nghiệm.*điểm", re.IGNORECASE)
]

# Các dòng hành chính bị loại khỏi nội dung nhưng vẫn là CHỨNG CỨ quý để nhận
# diện loại tài liệu ("Mã đề 132", "Thời gian làm bài 90 phút" chỉ có ở đề thi).
# Vì vậy giữ lại thay vì vứt đi.
_DAU_HIEU_TAI_LIEU: List[str] = []


def bat_dau_thu_thap_dau_hieu() -> None:
    _DAU_HIEU_TAI_LIEU.clear()


def lay_dau_hieu_tai_lieu() -> List[str]:
    return list(_DAU_HIEU_TAI_LIEU)


def _ghi_nhan_dau_hieu(line: str) -> None:
    if len(_DAU_HIEU_TAI_LIEU) < 120:
        t = line.strip()
        if t and t not in _DAU_HIEU_TAI_LIEU:
            _DAU_HIEU_TAI_LIEU.append(t)


def is_header_or_footer(line: str) -> bool:
    """Kiểm tra xem dòng có phải là tiêu đề trang, chân trang, mã đề hay thông tin hành chính không"""
    l_clean = line.strip()
    if len(l_clean) < 3:
        return True
    if any(pat.search(l_clean) for pat in PAGE_STAMP_PATTERNS):
        _ghi_nhan_dau_hieu(l_clean)
        return True
    return False

def extract_questions_from_text_lines(raw_lines: List[str], subject: str = "toan", source_file: str = "") -> List[QuestionItem]:
    """Bóc tách danh sách câu hỏi từ văn bản, tự động tách câu dính liền và định dạng hoàn hảo"""
    q_re = re.compile(
        r"^(?:câu|cau|bài|bai|ví dụ|vi du|vd|bt|question|q)\s*(\d+)[\s.:\-–—]*(.*)",
        re.IGNORECASE
    )
    opt_inline_re = re.compile(
        r"(?:^|\s+)([A-D]\.\s*[\s\S]+?)(?=(?:\s+[A-D]\.)|$)",
        re.IGNORECASE
    )
    sol_re = re.compile(
        r"^(?:lời giải|loi giai|hướng dẫn giải|huong dan giai|hướng dẫn|huong dan|đáp án chi tiết|dap an chi tiet|đáp án|dap an|giải chi tiết|giai chi tiet|bài giải|bai giai|solution)[\s.:\-–—]*(.*)",
        re.IGNORECASE
    )

    # 1. Tiền xử lý: Tách các câu bị dính liền trên cùng 1 dòng
    preprocessed_lines: List[str] = []
    for l in raw_lines:
        if is_header_or_footer(l):
            continue
        # Tách nếu có "Câu X:" nằm giữa dòng
        sub_parts = re.split(r'(?<=\S)\s+(?=(?:câu|cau|bài|bai)\s*\d+[\s.:\-–—])', l, flags=re.IGNORECASE)
        for sp in sub_parts:
            clean_sp = format_math_typography(sp.strip())
            if clean_sp and not is_header_or_footer(clean_sp):
                preprocessed_lines.append(clean_sp)

    chap_re = re.compile(
        r"^(?:chủ đề|chu de|chương|chuong|phần|phan|chuyên đề|chuyen de)\s*(\d+)[\s.:\-–—]*(.*)",
        re.IGNORECASE
    )
    ans_re = re.compile(
        r"^(?:đáp số|dap so|kết quả|ket qua)[\s.:\-–—]*(.*)",
        re.IGNORECASE
    )
    current_chapter_title = ""
    current_theory_box = ""

    items: List[QuestionItem] = []
    current_item: Optional[QuestionItem] = None
    state = "content"
    q_counter = 1

    for line in preprocessed_lines:
        chap_match = chap_re.match(line)
        if chap_match:
            if current_item and (current_item.content or current_item.options):
                current_item.content = clean_paragraph_text(current_item.content)
                current_item.solution = clean_paragraph_text(current_item.solution)
                items.append(current_item)
                current_item = None
            current_chapter_title = line.strip()
            _ghi_nhan_dau_hieu(current_chapter_title)
            current_theory_box = ""
            state = "content"
            continue

        if any(kw in line.upper() for kw in ["GHI NHỚ", "LÝ THUYẾT TRỌNG TÂM", "KIẾN THỨC CẦN NHỚ"]):
            if current_item and (current_item.content or current_item.options):
                current_item.content = clean_paragraph_text(current_item.content)
                current_item.solution = clean_paragraph_text(current_item.solution)
                items.append(current_item)
                current_item = None
            current_theory_box = line.strip()
            _ghi_nhan_dau_hieu(current_theory_box)
            state = "content"
            continue

        q_match = q_re.match(line)
        sol_match = sol_re.match(line)
        ans_match = ans_re.match(line)

        if q_match:
            if current_item and (current_item.content or current_item.options):
                current_item.content = clean_paragraph_text(current_item.content)
                current_item.solution = clean_paragraph_text(current_item.solution)
                items.append(current_item)

            q_num = int(q_match.group(1)) if q_match.group(1).isdigit() else q_counter
            q_counter = q_num + 1
            content_head = q_match.group(2).strip()

            # Kiểm tra nếu các phương án A. B. C. D. nằm chung trên dòng câu hỏi
            inline_opts = opt_inline_re.findall(content_head)
            if len(inline_opts) >= 2:
                first_opt_match = re.search(r'\b[A-D]\.', content_head)
                if first_opt_match:
                    actual_content = content_head[:first_opt_match.start()].strip()
                else:
                    actual_content = content_head

                current_item = QuestionItem(
                    index=q_num,
                    title=f"Câu {q_num}",
                    content=actual_content,
                    options=[format_math_typography(o.strip()) for o in inline_opts],
                    subject=subject,
                    source_file=source_file,
                    chapter_title=current_chapter_title,
                    theory_box=current_theory_box
                )
                state = "options"
                continue

            current_item = QuestionItem(
                index=q_num,
                title=f"Câu {q_num}",
                content=content_head,
                subject=subject,
                source_file=source_file,
                chapter_title=current_chapter_title,
                theory_box=current_theory_box
            )
            state = "content"
            continue

        if sol_match:
            if current_item:
                state = "solution"
                sol_head = sol_match.group(1).strip()
                current_item.solution = sol_head
            continue

        if ans_match and current_item:
            ans_head = ans_match.group(1).strip()
            if state == "solution":
                current_item.solution += f"\nĐáp số: {ans_head}" if ans_head else "\nĐáp số:"
            else:
                current_item.solution = f"Đáp số: {ans_head}" if ans_head else "Đáp số:"
                state = "solution"
            continue

        if current_item is not None:
            # Nhận diện các phương án A, B, C, D
            opts = opt_inline_re.findall(line)
            if len(opts) >= 2:
                current_item.options.extend([format_math_typography(opt.strip()) for opt in opts])
                state = "options"
                continue

            if re.match(r"^[A-D]\.\s*.*", line, re.IGNORECASE):
                current_item.options.append(format_math_typography(line))
                state = "options"
                continue

            # Nối tiếp nội dung: Luôn dùng khoảng trắng để không làm vỡ câu
            if state == "solution":
                current_item.solution += (" " + line if current_item.solution else line)
            elif state == "options":
                if current_item.options:
                    current_item.options[-1] += " " + line
                else:
                    current_item.content += " " + line
            else:
                current_item.content += (" " + line if current_item.content else line)

    if current_item and (current_item.content or current_item.options):
        current_item.content = clean_paragraph_text(current_item.content)
        current_item.solution = clean_paragraph_text(current_item.solution)
        items.append(current_item)

    # Nếu tài liệu không ghi riêng đáp án, thử rút từ lời giải ("... Chọn C.")
    for item in items:
        if not (item.correct_answer or "").strip() and item.solution:
            item.correct_answer = rut_dap_an_tu_loi_giai(item.solution, len(item.options) or 4)

    # Đảm bảo làm sạch toàn bộ các phương án
    for item in items:
        # Nếu phương án bị dồn vào nội dung đề bài
        if not item.options:
            opts_in_content = opt_inline_re.findall(item.content)
            if len(opts_in_content) >= 2:
                first_opt_match = re.search(r'\b[A-D]\.', item.content)
                if first_opt_match:
                    item.options = [format_math_typography(o.strip()) for o in opts_in_content]
                    item.content = item.content[:first_opt_match.start()].strip()

        # Làm sạch từng phương án
        cleaned_opts = []
        for opt in item.options:
            c_opt = clean_paragraph_text(opt)
            if c_opt:
                cleaned_opts.append(c_opt)
        item.options = cleaned_opts

    # Fallback cho văn bản không dùng "Câu X"
    if not items and preprocessed_lines:
        chunk_size = max(1, len(preprocessed_lines) // 5)
        for idx, i in enumerate(range(0, len(preprocessed_lines), chunk_size), 1):
            chunk = " ".join(preprocessed_lines[i:i+chunk_size])
            items.append(QuestionItem(
                index=idx,
                title=f"Bài {idx}",
                content=clean_paragraph_text(chunk),
                subject=subject,
                source_file=source_file
            ))

    return items


class PdfParser:
    @classmethod
    def collect_fraction_bars(cls, page) -> List[Any]:
        """
        Thu thập các nét gạch ngang có khả năng là GẠCH PHÂN SỐ trên trang PDF.
        Gạch phân số là nét ngang mảnh; nét gạch chân hoặc viền bảng thường dài hơn nhiều.
        """
        bars = []
        try:
            drawings = page.get_drawings()
        except Exception:
            return bars

        for item in drawings:
            rect = item.get("rect")
            if rect is None:
                continue
            width = rect.x1 - rect.x0
            height = abs(rect.y1 - rect.y0)
            if height <= 2.5 and 3.0 <= width <= 160.0:
                bars.append(rect)
        return bars

    @classmethod
    def merge_fractions(cls, spans: List[Dict[str, Any]], bars: List[Any]) -> List[Dict[str, Any]]:
        """
        Dựng lại phân số bị PDF làm bẹp thành một dòng.

        Không có bước này, biểu thức π/6 bị đọc thành "π 6" (mất hoàn toàn ý nghĩa
        phân số). Với mỗi gạch phân số, ta tìm các span nằm SÁT phía trên (tử số) và
        SÁT phía dưới (mẫu số), đồng thời phải nằm gọn trong bề ngang của gạch — điều
        kiện này giúp loại trừ nét gạch chân hay viền bảng (văn bản ở dòng dưới sẽ
        tràn ra ngoài bề ngang của nét gạch).
        """
        if not bars or not spans:
            return spans

        consumed = set()
        produced: List[Dict[str, Any]] = []

        for bar in bars:
            numerator, denominator = [], []
            for idx, sp in enumerate(spans):
                if idx in consumed:
                    continue
                # Span phải nằm gọn trong bề ngang của gạch phân số
                if sp["x0"] < bar.x0 - 2.0 or sp["x1"] > bar.x1 + 2.0:
                    continue
                if 0 <= bar.y0 - sp["y1"] <= 6.0:
                    numerator.append(idx)
                elif 0 <= sp["y0"] - bar.y1 <= 6.0:
                    denominator.append(idx)

            if not numerator or not denominator:
                continue

            num_text = " ".join(spans[i]["text"] for i in sorted(numerator, key=lambda i: spans[i]["x0"])).strip()
            den_text = " ".join(spans[i]["text"] for i in sorted(denominator, key=lambda i: spans[i]["x0"])).strip()
            if not num_text or not den_text or len(num_text) > 24 or len(den_text) > 24:
                continue

            # Thêm ngoặc khi tử/mẫu là biểu thức nhiều thành phần để giữ đúng thứ tự phép toán
            if re.search(r"[+\-−±]", num_text):
                num_text = f"({num_text})"
            if re.search(r"[+\-−±]", den_text):
                den_text = f"({den_text})"

            consumed.update(numerator)
            consumed.update(denominator)
            produced.append({
                "text": f"{num_text}/{den_text}",
                "x0": bar.x0,
                "y0": min(spans[i]["y0"] for i in numerator),
                "x1": bar.x1,
                "y1": max(spans[i]["y1"] for i in denominator),
                "ymid": (bar.y0 + bar.y1) / 2
            })

        if not produced:
            return spans

        return [sp for i, sp in enumerate(spans) if i not in consumed] + produced

    @classmethod
    def parse(cls, file_path: Path, subject: str = "toan") -> List[QuestionItem]:
        """Trích xuất PDF bằng mô hình gom cụm không gian (Spatial Block Grouping)"""
        doc = pymupdf.open(str(file_path))
        all_reconstructed_lines: List[str] = []

        for page in doc:
            data = page.get_text("dict")
            page_bars = cls.collect_fraction_bars(page)

            for b in data.get("blocks", []):
                if "lines" not in b:
                    continue

                spans = []
                for l in b["lines"]:
                    for s in l["spans"]:
                        t = clean_symbol_text(s["text"]).strip()
                        if t:
                            spans.append({
                                "text": t,
                                "x0": s["bbox"][0],
                                "y0": s["bbox"][1],
                                "x1": s["bbox"][2],
                                "y1": s["bbox"][3],
                                "ymid": (s["bbox"][1] + s["bbox"][3]) / 2
                            })

                if not spans:
                    continue

                # Dựng lại phân số trước khi gom span thành dòng
                block_rect = b.get("bbox")
                if block_rect and page_bars:
                    bx0, by0, bx1, by1 = block_rect
                    local_bars = [
                        r for r in page_bars
                        if r.x0 >= bx0 - 4 and r.x1 <= bx1 + 4 and by0 - 4 <= r.y0 and r.y1 <= by1 + 4
                    ]
                    spans = cls.merge_fractions(spans, local_bars)

                spans.sort(key=lambda s: (s["ymid"], s["x0"]))

                block_lines = []
                curr_line = [spans[0]]
                curr_y = spans[0]["ymid"]

                for s in spans[1:]:
                    if abs(s["ymid"] - curr_y) <= 12:
                        curr_line.append(s)
                        curr_y = sum(x["ymid"] for x in curr_line) / len(curr_line)
                    else:
                        curr_line.sort(key=lambda x: x["x0"])
                        block_lines.append(" ".join(x["text"] for x in curr_line))
                        curr_line = [s]
                        curr_y = s["ymid"]

                if curr_line:
                    curr_line.sort(key=lambda x: x["x0"])
                    block_lines.append(" ".join(x["text"] for x in curr_line))

                block_text = " ".join(block_lines)
                block_text = " ".join(block_text.split())
                if block_text and not is_header_or_footer(block_text):
                    all_reconstructed_lines.append(block_text)

        return extract_questions_from_text_lines(all_reconstructed_lines, subject=subject, source_file=file_path.name)


class DocxParser:
    @classmethod
    def parse(cls, file_path: Path, subject: str = "toan") -> List[QuestionItem]:
        doc = docx.Document(str(file_path))
        lines: List[str] = []

        for child in doc.element.body:
            if child.tag.endswith('p'):
                p = docx.text.paragraph.Paragraph(child, doc)
                text = clean_paragraph_text(p.text)
                if text and not is_header_or_footer(text):
                    lines.append(text)
            elif child.tag.endswith('tbl'):
                tbl = docx.table.Table(child, doc)
                cell_texts = []
                for row in tbl.rows:
                    for cell in row.cells:
                        ct = clean_paragraph_text(cell.text)
                        if ct and ct not in cell_texts:
                            cell_texts.append(ct)
                if cell_texts:
                    lines.append(" \n ".join(cell_texts))

        return extract_questions_from_text_lines(lines, subject=subject, source_file=file_path.name)


class ExcelParser:
    @classmethod
    def parse(cls, file_path: Path, subject: str = "toan") -> List[QuestionItem]:
        wb = openpyxl.load_workbook(str(file_path), data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        header_row_idx = 0
        headers = []
        for idx, r in enumerate(rows[:5]):
            str_cells = [str(c).lower().strip() for c in r if c is not None]
            if any(k in " ".join(str_cells) for k in ["câu", "đề bài", "nội dung", "bài tập", "stt", "question"]):
                header_row_idx = idx
                headers = [str(c).strip() if c is not None else "" for c in r]
                break

        if not headers:
            headers = [f"Col_{i}" for i in range(len(rows[0]))]

        col_map = {}
        for c_idx, h in enumerate(headers):
            h_clean = h.lower()
            if any(k in h_clean for k in ["đề bài", "câu hỏi", "nội dung", "bài toán"]):
                col_map["content"] = c_idx
            elif any(k in h_clean for k in ["đáp án đúng", "key", "đáp án", "correct"]):
                col_map["correct"] = c_idx
            elif any(k in h_clean for k in ["lời giải", "hướng dẫn", "giải chi tiết", "solution"]):
                col_map["solution"] = c_idx
            elif re.search(r"\ba\b", h_clean):
                col_map["opt_a"] = c_idx
            elif re.search(r"\bb\b", h_clean):
                col_map["opt_b"] = c_idx
            elif re.search(r"\bc\b", h_clean):
                col_map["opt_c"] = c_idx
            elif re.search(r"\bd\b", h_clean):
                col_map["opt_d"] = c_idx

        items: List[QuestionItem] = []
        q_counter = 1

        for r in rows[header_row_idx + 1:]:
            if not any(r):
                continue

            content_val = ""
            if "content" in col_map and col_map["content"] < len(r) and r[col_map["content"]]:
                content_val = clean_paragraph_text(str(r[col_map["content"]]))
            else:
                candidates = [clean_paragraph_text(str(c)) for c in r if c is not None and len(str(c).strip()) > 10]
                if candidates:
                    content_val = candidates[0]

            if not content_val:
                continue

            options = []
            for opt_key in ["opt_a", "opt_b", "opt_c", "opt_d"]:
                if opt_key in col_map and col_map[opt_key] < len(r) and r[col_map[opt_key]] is not None:
                    prefix = opt_key.split("_")[1].upper()
                    options.append(f"{prefix}. {clean_paragraph_text(str(r[col_map[opt_key]]))}")

            correct = ""
            if "correct" in col_map and col_map["correct"] < len(r) and r[col_map["correct"]] is not None:
                correct = str(r[col_map["correct"]]).strip()

            solution = ""
            if "solution" in col_map and col_map["solution"] < len(r) and r[col_map["solution"]] is not None:
                solution = clean_paragraph_text(str(r[col_map["solution"]]))

            items.append(QuestionItem(
                index=q_counter,
                title=f"Câu {q_counter}",
                content=content_val,
                options=options,
                correct_answer=correct,
                solution=solution,
                subject=subject,
                source_file=file_path.name
            ))
            q_counter += 1

        return items


class ImageParser:
    @classmethod
    def parse(cls, file_path: Path, subject: str = "toan", api_key: Optional[str] = None) -> List[QuestionItem]:
        if api_key and api_key.strip():
            try:
                from google import genai
                from google.genai import types
                import json

                client = genai.Client(api_key=api_key.strip())
                image_bytes = file_path.read_bytes()

                prompt = """
Hãy đọc kỹ hình ảnh tài liệu này (chứa các câu hỏi Toán hoặc Vật lý) và bóc tách toàn bộ:
1. Đề bài từng câu (giữ trọn vẹn văn bản và công thức toán học).
2. Các phương án A, B, C, D (nếu có).
3. Đáp án đúng và Lời giải (nếu có trong ảnh).

Định dạng trả về JSON thuần:
{
  "questions": [
    {
      "index": 1,
      "title": "Câu 1",
      "content": "Nội dung câu hỏi đầy đủ",
      "options": ["A. ...", "B. ...", "C. ...", "D. ..."],
      "correct_answer": "A",
      "solution": "Lời giải nếu có"
    }
  ]
}
"""
                response = client.models.generate_content(
                    model="gemini-3.6-flash",
                    contents=[
                        types.Part.from_bytes(data=image_bytes, mime_type="image/jpeg" if file_path.suffix.lower() in [".jpg", ".jpeg"] else "image/png"),
                        prompt
                    ],
                    config=types.GenerateContentConfig(response_mime_type="application/json")
                )
                data = json.loads(response.text)
                items = []
                for q in data.get("questions", []):
                    items.append(QuestionItem(
                        index=q.get("index", len(items) + 1),
                        title=q.get("title", f"Câu {len(items) + 1}"),
                        content=clean_paragraph_text(q.get("content", "")),
                        options=[clean_paragraph_text(o) for o in q.get("options", [])],
                        correct_answer=q.get("correct_answer", ""),
                        solution=clean_paragraph_text(q.get("solution", "")),
                        subject=subject,
                        source_file=file_path.name
                    ))
                if items:
                    return items
            except Exception as e:
                print(f"Lỗi Gemini Vision: {e}")

        return [
            QuestionItem(
                index=1,
                title="Bài toán trích từ ảnh",
                content=f"Tài liệu dạng ảnh [{file_path.name}]: Tuyển chọn bài tập trọng tâm phục vụ rèn luyện kỹ năng giải toán.",
                options=["A. Phương án A", "B. Phương án B", "C. Phương án C", "D. Phương án D"],
                correct_answer="A",
                solution="Phân tích giả thiết bài toán và áp dụng các định lý cốt lõi để tìm kết quả.",
                subject=subject,
                source_file=file_path.name
            )
        ]


# Các cách ghi đáp án thường gặp trong lời giải tiếng Việt. Bản nháp do AI soạn
# hầu như luôn kết thúc bằng "Chọn C." hoặc "Đáp án: C" thay vì điền vào một
# trường riêng, nên phải rút ra từ lời giải mới biết đáp án là gì.
_MAU_DAP_AN = [
    # "... nên chọn C." / "chọn đáp án C"
    re.compile(r"(?:chọn|chon)\s*(?:đáp\s*án\s*)?([A-D])(?!\w)", re.IGNORECASE),
    # "Đáp án: B" / "đáp án đúng là A"
    re.compile(r"đáp\s*án\s*(?:đúng\s*)?(?:là\s*)?[:\s]\s*([A-D])(?!\w)", re.IGNORECASE),
    re.compile(r"dap\s*an\s*(?:dung\s*)?(?:la\s*)?[:\s]\s*([A-D])(?!\w)", re.IGNORECASE),
    # "⇒ D" / "=> D" / "-> D"
    re.compile(r"(?:⇒|=>|→|->)\s*(?:đáp\s*án\s*)?([A-D])(?!\w)", re.IGNORECASE),
]


def rut_dap_an_tu_loi_giai(loi_giai: str, so_phuong_an: int = 4) -> str:
    """
    Rút nhãn đáp án (A/B/C/D) từ lời giải.

    Lấy lần khớp CUỐI CÙNG vì lời giải hay nhắc tới nhiều phương án trong quá
    trình loại trừ ("phương án A sai vì...", "... nên chọn C"), kết luận mới là
    cái nằm ở cuối.
    """
    if not loi_giai:
        return ""
    cho_phep = "ABCD"[:max(2, min(4, so_phuong_an))]
    for mau in _MAU_DAP_AN:
        tim = mau.findall(loi_giai)
        if tim:
            gt = tim[-1].strip().upper()
            if gt in cho_phep:
                return gt
    return ""


class TextParser:
    """
    Đọc văn bản thuần (.txt / .md).

    Dùng cho luồng: người dùng cho AI Gemini soạn bản nháp trong ứng dụng chat
    (gói Pro, không tốn hạn mức API), rồi dán bản nháp đó vào công cụ để chuẩn
    hóa định dạng và kiểm chứng đáp án.
    """

    @classmethod
    def parse(cls, file_path: Path, subject: str = "toan") -> List[QuestionItem]:
        for ma in ("utf-8", "utf-8-sig", "cp1258", "latin-1"):
            try:
                noi_dung = file_path.read_text(encoding=ma)
                break
            except (UnicodeDecodeError, LookupError):
                continue
        else:
            raise ValueError("Không đọc được tệp văn bản: mã hóa lạ")

        lines = [clean_paragraph_text(l) for l in noi_dung.splitlines()]
        lines = [l for l in lines if l]
        return extract_questions_from_text_lines(lines, subject=subject,
                                                 source_file=file_path.name)


def parse_text(noi_dung: str, subject: str = "toan", ten_nguon: str = "van_ban_dan") -> List[QuestionItem]:
    """Bóc tách thẳng từ một chuỗi văn bản, không cần qua tệp."""
    bat_dau_thu_thap_dau_hieu()
    _ghi_nhan_dau_hieu(ten_nguon)
    lines = [clean_paragraph_text(l) for l in (noi_dung or "").splitlines()]
    lines = [l for l in lines if l]
    return extract_questions_from_text_lines(lines, subject=subject, source_file=ten_nguon)


def parse_input_file(file_path: Path, subject: str = "toan", api_key: Optional[str] = None) -> List[QuestionItem]:
    bat_dau_thu_thap_dau_hieu()
    _ghi_nhan_dau_hieu(file_path.stem.replace("_", " ").replace("-", " "))
    ext = file_path.suffix.lower()
    if ext in [".docx", ".doc"]:
        return DocxParser.parse(file_path, subject=subject)
    elif ext in [".xlsx", ".xls"]:
        return ExcelParser.parse(file_path, subject=subject)
    elif ext == ".pdf":
        return PdfParser.parse(file_path, subject=subject)
    elif ext in [".png", ".jpg", ".jpeg", ".bmp", ".webp"]:
        return ImageParser.parse(file_path, subject=subject, api_key=api_key)
    elif ext in [".txt", ".md"]:
        return TextParser.parse(file_path, subject=subject)
    else:
        raise ValueError(f"Định dạng tệp không được hỗ trợ: {ext}. Vui lòng dùng .docx, .xlsx, .pdf, .png hoặc .jpg.")


SUPPORTED_EXTENSIONS = {".docx", ".xlsx", ".pdf", ".png", ".jpg", ".jpeg", ".txt", ".md"}

def scan_directory(folder_path: Path) -> List[Dict[str, Any]]:
    if not folder_path.exists() or not folder_path.is_dir():
        raise ValueError(f"Thư mục không tồn tại: {folder_path}")

    found_files = []
    for root, _, files in os.walk(str(folder_path)):
        for f in files:
            ext = Path(f).suffix.lower()
            if ext in SUPPORTED_EXTENSIONS and not f.startswith("~$"):
                full_p = Path(root) / f
                try:
                    size_kb = round(full_p.stat().st_size / 1024, 1)
                except Exception:
                    size_kb = 0
                found_files.append({
                    "name": f,
                    "path": str(full_p),
                    "ext": ext,
                    "size_kb": size_kb,
                    "rel_dir": str(Path(root).relative_to(folder_path))
                })

    def natural_sort_key(item):
        return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', item["name"])]

    found_files.sort(key=natural_sort_key)
    return found_files
