# -*- coding: utf-8 -*-
"""
Dựng ĐỀ KIỂM TRA / ĐỀ THI đúng cấu trúc ba phần áp dụng từ 2025.

Bộ đầy đủ gồm bốn phần, không phải một tệp đề:
    1. Ma trận đề        — phân bố câu theo chuyên đề × mức độ
    2. Bản đặc tả        — mỗi câu kiểm tra yêu cầu cần đạt nào
    3. Đề kiểm tra       — bản phát cho học sinh, không lộ đáp án
    4. Hướng dẫn chấm    — đáp án kèm thang điểm từng phần

Ma trận và bản đặc tả là hai thứ giáo viên tốn nhiều công nhất, nhưng lại suy
ra được bằng máy vì mỗi câu đã mang sẵn mức độ và chuyên đề.

GIỚI HẠN ĐÃ BIẾT của Phần II: tài liệu gốc thường là trắc nghiệm bốn lựa chọn,
không đủ dữ kiện để dựng bốn ý Đúng/Sai thật sự. Ở đây chỉ làm phép biến đổi
CƠ HỌC và KIỂM CHỨNG ĐƯỢC: mỗi ý khẳng định "đáp án đúng là X", nên đúng một ý
Đúng và ba ý Sai — không bịa thêm dữ kiện nào. Bản đặc tả ghi rõ những câu như
vậy để giáo viên rà lại và viết thành ngữ cảnh chung cho hay hơn.
"""

from typing import List, Dict, Any, Tuple

import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from core.loai_dau_ra import (
    tinh_cau_truc, quy_cach, phan_bo_muc_do, kiem_tra_tong_diem, MUC_DO,
)

FONT_MAIN = "Times New Roman"
COLOR_PRIMARY = RGBColor(26, 54, 93)
COLOR_MUTED = RGBColor(74, 85, 104)
COLOR_WARN = RGBColor(197, 48, 48)

NHAN_Y = ("a", "b", "c", "d")


# ---------------------------------------------------------------------------
# Tiện ích dựng văn bản
# ---------------------------------------------------------------------------
def _p(doc, text="", *, bold=False, italic=False, size=12, align=None,
       color=None, before=0, after=0):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    if text:
        r = para.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.name = FONT_MAIN
        r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
    return para


def _tieu_de_phan(doc, text, ghi_chu=""):
    _p(doc, text, bold=True, size=13, color=COLOR_PRIMARY, before=14, after=2)
    if ghi_chu:
        _p(doc, ghi_chu, italic=True, size=11, color=COLOR_MUTED, after=6)


def _bang(doc, tieu_de: List[str], dong: List[List[str]], rong: List[int] = None):
    t = doc.add_table(rows=1, cols=len(tieu_de))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, nhan in enumerate(tieu_de):
        o = t.rows[0].cells[i]
        o.text = ""
        r = o.paragraphs[0].add_run(nhan)
        r.bold = True
        r.font.name = FONT_MAIN
        r.font.size = Pt(10.5)
        o.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for d in dong:
        cells = t.add_row().cells
        for i, gt in enumerate(d):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(gt))
            r.font.name = FONT_MAIN
            r.font.size = Pt(10.5)
            if i > 0:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return t


def _so(x: float) -> str:
    """Việt Nam dùng dấu phẩy thập phân: 0,25 chứ không phải 0.25."""
    s = f"{x:.3f}".rstrip("0").rstrip(".")
    return (s or "0").replace(".", ",")


# ---------------------------------------------------------------------------
# Chia câu hỏi về ba phần
# ---------------------------------------------------------------------------
def _la_dap_so(q) -> bool:
    """Câu hợp với Phần III là câu mà đáp án vốn đã là một con số."""
    import re
    dap = (getattr(q, "correct_answer", "") or "").strip().upper()[:1]
    opts = getattr(q, "new_options", None) or []
    for o in opts:
        nhan = str(o).strip()[:1].upper()
        if nhan == dap:
            than = str(o).strip()[2:].strip()
            return bool(re.fullmatch(r"[-+]?\d+([.,]\d+)?\s*[^\s]{0,8}", than))
    return False


def chia_ba_phan(questions: List[Any], ct) -> Tuple[list, list, list]:
    """
    Xếp câu vào ba phần.

    Phần III ưu tiên những câu vốn đã có đáp số là một con số — đúng yêu cầu
    "đáp án phải là một số điền được vào ô" của Bộ. Nếu không đủ thì phần thiếu
    để trống, chứ KHÔNG ép một câu có đáp án là mệnh đề thành câu trả lời ngắn.
    """
    con_lai = list(questions)

    hop_p3 = [q for q in con_lai if _la_dap_so(q)]
    p3 = hop_p3[: ct.so_cau_p3]
    dung_roi = {id(q) for q in p3}
    con_lai = [q for q in con_lai if id(q) not in dung_roi]

    p1 = con_lai[: ct.so_cau_p1]
    con_lai = con_lai[ct.so_cau_p1:]

    p2 = con_lai[: ct.so_cau_p2]
    return p1, p2, p3


def y_dung_sai(q) -> List[Tuple[str, bool]]:
    """
    Biến một câu trắc nghiệm thành bốn ý Đúng/Sai một cách CƠ HỌC.

    Mỗi ý khẳng định một phương án là đáp án đúng. Đúng một ý Đúng, ba ý Sai —
    không bịa thêm dữ kiện nào, nên luôn kiểm chứng được. Chất lượng sư phạm
    thì chưa bằng ngữ cảnh chung do giáo viên viết; bản đặc tả có ghi chú.
    """
    dap = (getattr(q, "correct_answer", "") or "").strip().upper()[:1]
    ds = []
    for o in (getattr(q, "new_options", None) or [])[:4]:
        s = str(o).strip()
        nhan = s[:1].upper()
        than = s[2:].strip() if len(s) > 2 else s
        ds.append((f"Kết quả của bài toán là {than}", nhan == dap))
    while len(ds) < 4:
        ds.append(("(Giáo viên bổ sung ý này)", False))
    return ds[:4]


def diem_phan_ii(so_y_dung: int, ct) -> float:
    """
    Thang lũy tiến: KHÔNG phải 0,25 nhân số ý đúng.
    Đúng 2 trong 4 ý chỉ được 0,25 chứ không phải 0,5.
    """
    if so_y_dung <= 0:
        return 0.0
    return ct.luy_tien_p2[min(so_y_dung, 4) - 1]


# ---------------------------------------------------------------------------
# 1. MA TRẬN ĐỀ
# ---------------------------------------------------------------------------
def render_ma_tran(doc, book, ma_loai: str, subject: str, p1, p2, p3, ct):
    q = quy_cach(ma_loai)
    _p(doc, "MA TRẬN ĐỀ KIỂM TRA", bold=True, size=15, color=COLOR_PRIMARY,
       align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    _p(doc, f"{q.ten} — Môn {'Toán' if subject == 'toan' else 'Vật lí'}"
            f" — Thời gian {ct.thoi_gian} phút",
       italic=True, size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    _p(doc, f"Phạm vi kiến thức: {q.pham_vi}", size=11.5, after=10)

    # Bảng 1: cấu trúc và thang điểm
    _tieu_de_phan(doc, "1. Cấu trúc đề và thang điểm")
    _bang(doc,
          ["Phần", "Dạng thức", "Số câu", "Số lệnh hỏi", "Điểm mỗi câu", "Tổng điểm"],
          [
              ["I", "Trắc nghiệm nhiều lựa chọn", ct.so_cau_p1, ct.so_cau_p1,
               _so(ct.diem_moi_cau_p1), _so(ct.so_cau_p1 * ct.diem_moi_cau_p1)],
              ["II", "Trắc nghiệm Đúng/Sai (4 ý)", ct.so_cau_p2, ct.so_cau_p2 * 4,
               f"tối đa {_so(ct.diem_toi_da_p2)}", _so(ct.so_cau_p2 * ct.diem_toi_da_p2)],
              ["III", "Trả lời ngắn", ct.so_cau_p3, ct.so_cau_p3,
               _so(ct.diem_moi_cau_p3), _so(ct.so_cau_p3 * ct.diem_moi_cau_p3)],
              ["", "TỔNG", ct.tong_cau, ct.tong_lenh_hoi, "", _so(ct.tong_diem)],
          ])

    if not kiem_tra_tong_diem(ct):
        _p(doc, f"⚠ Tổng điểm đang là {_so(ct.tong_diem)}, chưa bằng 10,0 — cần rà lại.",
           bold=True, size=11, color=COLOR_WARN, before=4)

    _p(doc, "Thang điểm Phần II tính lũy tiến theo số ý đúng trong mỗi câu: "
            f"1 ý {_so(ct.luy_tien_p2[0])} điểm · 2 ý {_so(ct.luy_tien_p2[1])} điểm · "
            f"3 ý {_so(ct.luy_tien_p2[2])} điểm · 4 ý {_so(ct.luy_tien_p2[3])} điểm. "
            "Đây KHÔNG phải phép nhân đều — đúng 2 trong 4 ý chỉ được "
            f"{_so(ct.luy_tien_p2[1])} điểm.",
       italic=True, size=11, color=COLOR_MUTED, before=6, after=4)

    if q.thoi_gian and ct.tong_cau < 22:
        _p(doc, "Ghi chú: đây là bài kiểm tra định kì trong trường nên số câu và "
                "thời gian do nhà trường quyết định theo Thông tư 22/2021. Điểm mỗi "
                "câu vì thế khác đề tốt nghiệp, nhưng tỉ trọng giữa ba phần và "
                "thang lũy tiến Phần II thì giữ nguyên.",
           italic=True, size=10.5, color=COLOR_MUTED, after=6)

    # Bảng 2: phân bố theo chuyên đề × mức độ
    _tieu_de_phan(doc, "2. Phân bố theo nội dung và mức độ tư duy")
    theo_chu_de: Dict[str, Dict[str, int]] = {}
    for phan, ds in (("I", p1), ("II", p2), ("III", p3)):
        for q_ in ds:
            cd = _chu_de_cua(q_)
            muc = _muc_do_cua(q_)
            theo_chu_de.setdefault(cd, {m: 0 for m in MUC_DO})
            so_lenh = 4 if phan == "II" else 1
            theo_chu_de[cd][muc] += so_lenh

    dong = []
    tong = {m: 0 for m in MUC_DO}
    for cd, d in theo_chu_de.items():
        dong.append([cd] + [d[m] for m in MUC_DO] + [sum(d.values())])
        for m in MUC_DO:
            tong[m] += d[m]
    dong.append(["TỔNG"] + [tong[m] for m in MUC_DO] + [sum(tong.values())])
    _bang(doc, ["Nội dung / Chuyên đề"] + list(MUC_DO) + ["Tổng lệnh hỏi"], dong)

    goi_y = phan_bo_muc_do(ct.tong_lenh_hoi)
    _p(doc, "Tỉ lệ tham khảo của Bộ cho ba mức: 40% Biết – 30% Hiểu – 30% Vận dụng, "
            f"ứng với đề này là {goi_y['Biết']} – {goi_y['Hiểu']} – {goi_y['Vận dụng']} lệnh hỏi.",
       italic=True, size=11, color=COLOR_MUTED, before=6)


_CACHE_CHU_DE: Dict[int, str] = {}


def _chu_de_cua(q) -> str:
    """
    Tên chuyên đề của một câu, dùng để dựng ma trận.

    Nhận diện theo nội dung câu hỏi bằng ngân hàng chuyên đề có sẵn. Không nhận
    ra thì xếp vào một nhóm chung, chứ không bỏ câu đó khỏi ma trận — ma trận
    thiếu câu thì tổng không khớp và giáo viên không dùng được.
    """
    khoa = id(q)
    if khoa in _CACHE_CHU_DE:
        return _CACHE_CHU_DE[khoa]

    ten = "Nội dung chung"
    try:
        from core.theory_bank import (
            detect_subject_and_topic, MATH_THEORY_DATABASE, PHYSICS_THEORY_DATABASE)
        nd = (getattr(q, "new_content", "") or getattr(q, "original_content", "") or "")
        kq = detect_subject_and_topic([nd])
        if kq.get("match_score", 0) > 0:
            khoa_cd = kq.get("topic_key", "")
            muc = (MATH_THEORY_DATABASE.get(khoa_cd)
                   or PHYSICS_THEORY_DATABASE.get(khoa_cd) or {})
            tieu_de = muc.get("title") or ""
            # Bỏ tiền tố "CHUYÊN ĐỀ: " cho ô bảng khỏi dài lê thê
            ten = tieu_de.split(":", 1)[-1].strip() if ":" in tieu_de else (
                tieu_de or khoa_cd.replace("_", " ").title() or ten)
    except Exception:
        pass

    _CACHE_CHU_DE[khoa] = ten
    return ten


def _muc_do_cua(q) -> str:
    """Quy mọi cách ghi mức độ về đúng ba mức của chuẩn 2025."""
    lv = (getattr(q, "level", "") or "").strip().lower()
    if "vận dụng" in lv or "van dung" in lv:
        return "Vận dụng"
    if "hiểu" in lv or "hieu" in lv:
        return "Hiểu"
    return "Biết"


# ---------------------------------------------------------------------------
# 2. BẢN ĐẶC TẢ
# ---------------------------------------------------------------------------
def render_dac_ta(doc, ma_loai, subject, p1, p2, p3, ct):
    doc.add_page_break()
    _p(doc, "BẢN ĐẶC TẢ ĐỀ KIỂM TRA", bold=True, size=15, color=COLOR_PRIMARY,
       align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    _p(doc, "Mỗi câu kiểm tra yêu cầu cần đạt nào, ở mức độ tư duy nào.",
       italic=True, size=11, color=COLOR_MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

    dong = []
    can_ra_lai = 0
    for phan, ds, diem in (("I", p1, ct.diem_moi_cau_p1),
                           ("II", p2, ct.diem_toi_da_p2),
                           ("III", p3, ct.diem_moi_cau_p3)):
        for i, q in enumerate(ds, 1):
            ghi = ""
            if phan == "II":
                ghi = "Chuyển cơ học từ trắc nghiệm — nên rà lại"
                can_ra_lai += 1
            noi_dung = (getattr(q, "new_content", "") or "")[:90]
            dong.append([f"{phan}.{i}", _chu_de_cua(q), _muc_do_cua(q),
                         _so(diem), noi_dung + ("…" if len(noi_dung) >= 90 else ""), ghi])

    _bang(doc, ["Câu", "Chuyên đề", "Mức độ", "Điểm", "Yêu cầu cần đạt", "Ghi chú"], dong)

    if can_ra_lai:
        _p(doc, f"⚠ {can_ra_lai} câu ở Phần II được chuyển cơ học từ câu trắc nghiệm "
                "bốn lựa chọn: mỗi ý khẳng định một phương án là đáp án đúng, nên đúng "
                "một ý Đúng và ba ý Sai. Cách này luôn kiểm chứng được và không bịa "
                "dữ kiện, nhưng chưa bằng một ngữ cảnh chung do giáo viên viết. "
                "Nên đọc lại và viết lại bốn ý quanh cùng một tình huống.",
           size=11, color=COLOR_WARN, before=8)


# ---------------------------------------------------------------------------
# 3. ĐỀ KIỂM TRA
# ---------------------------------------------------------------------------
def render_de(doc, book, ma_loai, subject, p1, p2, p3, ct):
    doc.add_page_break()
    q_loai = quy_cach(ma_loai)
    mon = "TOÁN" if subject == "toan" else "VẬT LÍ"

    thong_tin = getattr(book, "exam_info", {}) or {}
    _p(doc, thong_tin.get("truong", "SỞ GIÁO DỤC VÀ ĐÀO TẠO ………………"),
       bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, q_loai.ten.upper(), bold=True, size=14, color=COLOR_PRIMARY,
       align=WD_ALIGN_PARAGRAPH.CENTER, before=4)
    _p(doc, f"Môn: {mon}", bold=True, size=12.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, f"Thời gian làm bài: {ct.thoi_gian} phút, không kể thời gian phát đề",
       italic=True, size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    _p(doc, "Họ và tên học sinh: ……………………………………………  Lớp: …………  Mã đề: ………",
       size=12, after=10)

    # PHẦN I
    _tieu_de_phan(
        doc, "PHẦN I. TRẮC NGHIỆM NHIỀU LỰA CHỌN",
        f"Học sinh trả lời từ câu 1 đến câu {len(p1)}. "
        f"Mỗi câu chỉ chọn MỘT phương án. Mỗi câu đúng được {_so(ct.diem_moi_cau_p1)} điểm.")
    for i, q in enumerate(p1, 1):
        _p(doc, f"Câu {i}. {getattr(q, 'new_content', '')}", size=12, before=6, after=2)
        for o in (getattr(q, "new_options", None) or []):
            _p(doc, f"    {o}", size=12, after=0)

    # PHẦN II
    _tieu_de_phan(
        doc, "PHẦN II. TRẮC NGHIỆM ĐÚNG/SAI",
        f"Học sinh trả lời từ câu 1 đến câu {len(p2)}. Trong mỗi ý a), b), c), d) "
        f"ở mỗi câu, học sinh chọn ĐÚNG hoặc SAI. "
        f"Điểm tính lũy tiến theo số ý đúng: {_so(ct.luy_tien_p2[0])} – "
        f"{_so(ct.luy_tien_p2[1])} – {_so(ct.luy_tien_p2[2])} – {_so(ct.luy_tien_p2[3])}.")
    for i, q in enumerate(p2, 1):
        _p(doc, f"Câu {i}. {getattr(q, 'new_content', '')}", size=12, before=6, after=2)
        for nhan, (noi_dung, _) in zip(NHAN_Y, y_dung_sai(q)):
            _p(doc, f"    {nhan}) {noi_dung}", size=12, after=0)

    # PHẦN III
    _tieu_de_phan(
        doc, "PHẦN III. TRẢ LỜI NGẮN",
        f"Học sinh trả lời từ câu 1 đến câu {len(p3)}. Viết đáp số vào ô trả lời. "
        f"Mỗi câu đúng được {_so(ct.diem_moi_cau_p3)} điểm.")
    for i, q in enumerate(p3, 1):
        _p(doc, f"Câu {i}. {getattr(q, 'new_content', '')}", size=12, before=6, after=2)
        _p(doc, "    Đáp số: ……………………", size=12, after=0)

    _p(doc, "---------- HẾT ----------", bold=True, size=12,
       align=WD_ALIGN_PARAGRAPH.CENTER, before=16)
    _p(doc, "Học sinh không được sử dụng tài liệu. "
            "Cán bộ coi thi không giải thích gì thêm.",
       italic=True, size=10.5, color=COLOR_MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)


# ---------------------------------------------------------------------------
# 4. HƯỚNG DẪN CHẤM
# ---------------------------------------------------------------------------
def render_huong_dan_cham(doc, p1, p2, p3, ct):
    doc.add_page_break()
    _p(doc, "HƯỚNG DẪN CHẤM", bold=True, size=15, color=COLOR_PRIMARY,
       align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

    # Phần I
    _tieu_de_phan(doc, f"PHẦN I — mỗi câu {_so(ct.diem_moi_cau_p1)} điểm")
    if p1:
        moi_hang = 10
        for dau in range(0, len(p1), moi_hang):
            lat = p1[dau:dau + moi_hang]
            _bang(doc,
                  [f"Câu {dau + i + 1}" for i in range(len(lat))],
                  [[(getattr(q, "correct_answer", "") or "?").upper()[:1] for q in lat]])

    # Phần II
    _tieu_de_phan(doc, f"PHẦN II — lũy tiến {_so(ct.luy_tien_p2[0])} / "
                       f"{_so(ct.luy_tien_p2[1])} / {_so(ct.luy_tien_p2[2])} / "
                       f"{_so(ct.luy_tien_p2[3])} điểm theo số ý đúng")
    if p2:
        dong = []
        for i, q in enumerate(p2, 1):
            ds = y_dung_sai(q)
            dong.append([f"Câu {i}"] + ["Đúng" if dung else "Sai" for _, dung in ds])
        _bang(doc, ["Câu", "a)", "b)", "c)", "d)"], dong)

    # Phần III
    _tieu_de_phan(doc, f"PHẦN III — mỗi câu {_so(ct.diem_moi_cau_p3)} điểm")
    if p3:
        dong = []
        for i, q in enumerate(p3, 1):
            dap = (getattr(q, "correct_answer", "") or "").upper()[:1]
            than = ""
            for o in (getattr(q, "new_options", None) or []):
                if str(o).strip()[:1].upper() == dap:
                    than = str(o).strip()[2:].strip()
            dong.append([f"Câu {i}", than or "…"])
        _bang(doc, ["Câu", "Đáp số"], dong)

    _p(doc, f"Tổng điểm toàn bài: {_so(ct.tong_diem)}", bold=True, size=12, before=12)


# ---------------------------------------------------------------------------
# Điểm vào
# ---------------------------------------------------------------------------
def xuat_bo_de(doc, book, ma_loai: str, subject: str = "toan",
               kem_loi_giai: bool = True):
    """Dựng trọn bộ bốn phần vào tài liệu Word đang mở."""
    ds_cau = list(getattr(book, "questions", None) or [])
    if not ds_cau:
        for ch in getattr(book, "chapters", None) or []:
            ds_cau.extend(ch.questions)

    ct = tinh_cau_truc(ma_loai, subject)
    p1, p2, p3 = chia_ba_phan(ds_cau, ct)

    render_ma_tran(doc, book, ma_loai, subject, p1, p2, p3, ct)
    render_dac_ta(doc, ma_loai, subject, p1, p2, p3, ct)
    render_de(doc, book, ma_loai, subject, p1, p2, p3, ct)
    render_huong_dan_cham(doc, p1, p2, p3, ct)

    if kem_loi_giai:
        doc.add_page_break()
        _p(doc, "HƯỚNG DẪN GIẢI CHI TIẾT", bold=True, size=15, color=COLOR_PRIMARY,
           align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
        for phan, ds in (("I", p1), ("II", p2), ("III", p3)):
            for i, q in enumerate(ds, 1):
                _p(doc, f"Phần {phan} — Câu {i}", bold=True, size=12,
                   color=COLOR_PRIMARY, before=10, after=2)
                lg = (getattr(q, "solution_method1", "") or "").strip()
                if lg:
                    _p(doc, lg, size=12, after=2)
                lg2 = (getattr(q, "solution_method2", "") or "").strip()
                if lg2:
                    _p(doc, lg2, italic=True, size=11.5, color=COLOR_MUTED, after=2)

    return {
        "so_cau_p1": len(p1), "so_cau_p2": len(p2), "so_cau_p3": len(p3),
        "tong_diem": ct.tong_diem, "thoi_gian": ct.thoi_gian,
        "du_cau": len(p1) == ct.so_cau_p1 and len(p2) == ct.so_cau_p2
                  and len(p3) == ct.so_cau_p3,
    }
