# -*- coding: utf-8 -*-
"""
Dựng CHUYÊN ĐỀ BÀI TẬP theo khuôn bốn phần.

VÌ SAO CÓ TỆP NÀY. Trước đây `CHUYEN_DE_BT` không có nhánh rẽ riêng nên rơi vào
bộ dựng sách mặc định — mà bộ đó **in lời giải ngay dưới mỗi bài**. Skill
`chuyen-de-bai-tap`, đo từ video mẫu do chủ dự án cung cấp, quy định ngược lại:

    "Đáp án dồn hết về cuối tài liệu, không in ngay dưới mỗi câu.
     Học sinh phải tự làm trước rồi mới đối chiếu."

Bộ kiểm kỷ luật (`scripts/kiem_ky_luat.py`, luật KL4) đã bắt được chỗ này.

Bốn phần bắt buộc:

    Phần 1. Lý thuyết và Ví dụ minh họa      ví dụ có lời giải NGAY tại chỗ
    Phần 2. Bài tập tự luyện                 chia ba dải mức độ, KHÔNG có đáp án
    Phần 3. Câu hỏi từ đề thi                đánh số "Câu T1." kèm thẻ năm
    Phần 4. Hướng dẫn và Đáp án              bảng ba cột, dồn hết về cuối

MỘT ĐIỀU KHÔNG ĐƯỢC LÀM: Phần 3 chỉ nhận câu **thật sự truy được nguồn** trong
tài liệu gốc. Nếu tài liệu đầu vào không có câu nào ghi năm hay kỳ thi thì phần
này để trống kèm ghi chú, TUYỆT ĐỐI không tự gán thẻ `[2015]` cho một câu không
rõ xuất xứ — gán bừa là bịa nguồn, học sinh tin rằng câu đó từng ra thi thật.
"""

import re
from typing import Any, Dict, List, Tuple

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

FONT_MAIN = "Times New Roman"
COLOR_PRIMARY = RGBColor(26, 54, 93)
COLOR_MUTED = RGBColor(74, 85, 104)
COLOR_MUC_DO = RGBColor(192, 86, 33)     # cam, dùng cho dải mức độ

# Ba mức của CHUYÊN ĐỀ khác hẳn ba mức của đề kiểm tra 2025. Tài liệu luyện tập
# dùng Dễ / Trung bình / Khó; đề chuẩn 2025 dùng Biết / Hiểu / Vận dụng. Skill
# ghi rõ: không trộn hai thang này.
MUC_DE, MUC_TB, MUC_KHO = "DỄ", "TRUNG BÌNH", "KHÓ"
THU_TU_MUC = (MUC_DE, MUC_TB, MUC_KHO)

_QUY_MUC = {
    "nhận biết": MUC_DE, "nhan biet": MUC_DE, "biết": MUC_DE, "biet": MUC_DE,
    "thông hiểu": MUC_TB, "thong hieu": MUC_TB, "hiểu": MUC_TB, "hieu": MUC_TB,
    "vận dụng": MUC_KHO, "van dung": MUC_KHO,
    "vận dụng cao": MUC_KHO, "van dung cao": MUC_KHO,
}

# Dấu hiệu một câu có xuất xứ từ đề thi thật. Chỉ những câu khớp mới được vào
# Phần 3, và thẻ năm lấy đúng con số tìm thấy chứ không suy đoán.
_NAM = re.compile(r"(?:^|[\s\[\(])((?:19|20)\d{2})(?:\s*[-–]\s*(?:19|20)?\d{2})?(?=[\s\]\)\.,;:]|$)")
_DAU_HIEU_DE = re.compile(
    r"(đề\s*thi|tốt\s*nghiệp|thpt\s*qg|thpt\s*quốc\s*gia|tuyển\s*sinh|"
    r"minh\s*họa|tham\s*khảo|chính\s*thức|đề\s*minh)", re.IGNORECASE)


def _p(doc, text="", *, bold=False, italic=False, underline=False, size=12,
       align=None, color=None, before=0, after=0, thut=0.0):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if thut:
        pf.left_indent = Pt(thut)
    if text:
        r = para.add_run(text)
        r.bold = bold
        r.italic = italic
        r.underline = underline
        r.font.name = FONT_MAIN
        r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
    return para


def _tieu_de_phan(doc, text: str):
    _p(doc, text, bold=True, size=14, color=COLOR_PRIMARY, before=22, after=8)


def _muc_do_cua(q) -> str:
    return _QUY_MUC.get((getattr(q, "level", "") or "").strip().lower(), MUC_TB)


def _nam_cua(q) -> str:
    """Trả về năm nếu câu truy được nguồn đề thi; chuỗi rỗng nghĩa là không rõ."""
    van = " ".join(str(getattr(q, t, "") or "") for t in
                   ("title", "new_content", "original_content", "source_file"))
    if not _DAU_HIEU_DE.search(van):
        return ""
    m = _NAM.search(van)
    return m.group(1) if m else ""


def _tach_phan_2_va_3(ds: List[Any]) -> Tuple[List[Any], List[Tuple[Any, str]]]:
    """Câu truy được nguồn đề thi thì sang Phần 3, còn lại ở Phần 2."""
    tu_luyen, tu_de_thi = [], []
    for q in ds:
        nam = _nam_cua(q)
        if nam:
            tu_de_thi.append((q, nam))
        else:
            tu_luyen.append(q)
    tu_de_thi.sort(key=lambda x: x[1])
    return tu_luyen, tu_de_thi


def _cum_nam(nam: str) -> str:
    """Gom theo cụm 2-3 năm như mẫu, không liệt kê từng năm riêng lẻ."""
    try:
        n = int(nam)
    except ValueError:
        return "Chưa rõ năm"
    dau = n - (n % 3)
    return "NĂM %d–%d" % (dau, dau + 2)


def _phuong_an(q) -> List[str]:
    return [str(o).strip() for o in (getattr(q, "new_options", None) or []) if str(o).strip()]


def _in_cau(doc, nhan: str, q, cb_size: float):
    """In một câu KHÔNG kèm đáp án — đáp án dồn hết về Phần 4."""
    _p(doc, f"{nhan} {getattr(q, 'new_content', '') or ''}",
       size=cb_size, before=8, after=2)
    pa = _phuong_an(q)
    if pa:
        # Bốn phương án trên cùng một dòng cho gọn, như mẫu
        _p(doc, "   ".join(pa), size=cb_size, after=2, thut=18)


def _bang_dap_an(doc, tieu_de: str, cac_dong: List[Tuple[str, str, str]]):
    """Bảng ba cột Câu / Đáp án / Hướng dẫn giải — mẫu quy định KHÔNG viết văn xuôi."""
    _p(doc, tieu_de, bold=True, size=12.5, color=COLOR_PRIMARY, before=14, after=4)
    bang = doc.add_table(rows=1, cols=3)
    bang.style = "Table Grid"
    dau = bang.rows[0].cells
    for i, ten in enumerate(("Câu", "Đáp án", "Hướng dẫn giải")):
        dau[i].text = ""
        r = dau[i].paragraphs[0].add_run(ten)
        r.bold = True
        r.font.name = FONT_MAIN
        r.font.size = Pt(11.5)
        r.font.color.rgb = COLOR_PRIMARY

    for cau, dap, huong in cac_dong:
        o = bang.add_row().cells
        for i, v in enumerate((cau, dap, huong)):
            o[i].text = ""
            r = o[i].paragraphs[0].add_run(v)
            r.font.name = FONT_MAIN
            r.font.size = Pt(11)


def _tom_huong_dan(q, gioi_han: int = 260) -> str:
    """
    Cột hướng dẫn viết gọn — một tới ba dòng nêu đúng bước then chốt, không chép
    lại toàn bộ lời giải dài. Cắt ở ranh giới câu cho khỏi cụt giữa chừng.
    """
    lg = (getattr(q, "solution_method1", "") or "").strip()
    lg = re.sub(r"\s+", " ", lg)
    if len(lg) <= gioi_han:
        return lg
    cat = lg[:gioi_han]
    dau_cham = cat.rfind(". ")
    return (cat[:dau_cham + 1] if dau_cham > 60 else cat.rstrip()) + " …"


def xuat_chuyen_de(doc, book, subject: str = "toan", options: Dict[str, Any] = None):
    """Dựng trọn chuyên đề bốn phần vào tài liệu Word đang mở."""
    opts = options or {}
    ds: List[Any] = list(getattr(book, "questions", None) or [])
    if not ds:
        for ch in getattr(book, "chapters", None) or []:
            ds.extend(ch.questions)

    co_chu = 12.0
    mon = "Toán học" if subject == "toan" else "Vật lí"
    ten = (getattr(book, "new_title", "") or "CHUYÊN ĐỀ BÀI TẬP").strip()

    # ---------- Trang đầu ----------
    _p(doc, ten.upper(), bold=True, size=17, color=COLOR_PRIMARY,
       align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    _p(doc, "Lý thuyết — Bài tập — Câu hỏi từ đề thi", italic=True, size=12.5,
       align=WD_ALIGN_PARAGRAPH.CENTER, color=COLOR_MUTED, after=2)
    _p(doc, f"{mon}", size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER,
       color=COLOR_MUTED, after=16)

    tu_luyen, tu_de_thi = _tach_phan_2_va_3(ds)

    # ---------- Mục lục ----------
    _p(doc, "MỤC LỤC", bold=True, size=13, color=COLOR_PRIMARY, after=4)
    for dong in (
        "Phần 1. Lý thuyết và Ví dụ minh họa",
        f"Phần 2. Bài tập tự luyện ({len(tu_luyen)} câu, từ dễ đến khó)",
        f"Phần 3. Câu hỏi từ đề thi ({len(tu_de_thi)} câu)",
        "Phần 4. Hướng dẫn và Đáp án",
    ):
        _p(doc, "   " + dong, size=12, after=1)

    # ---------- PHẦN 1 ----------
    _tieu_de_phan(doc, "PHẦN 1. LÝ THUYẾT VÀ VÍ DỤ MINH HỌA")
    ly_thuyet = (getattr(book, "theory_section", "") or "").strip()
    if ly_thuyet:
        for doan in [d.strip() for d in ly_thuyet.split("\n") if d.strip()]:
            # Dòng ngắn không kết thúc bằng dấu câu thì coi là đầu mục
            la_dau_muc = len(doan) < 80 and not doan.endswith((".", ":", ";", ","))
            _p(doc, doan, bold=la_dau_muc, size=co_chu,
               color=COLOR_PRIMARY if la_dau_muc else None,
               before=8 if la_dau_muc else 0, after=3)
    else:
        _p(doc, "(Tài liệu gốc không kèm phần lý thuyết. Bổ sung trước khi đưa "
                "cho học sinh — chuyên đề thiếu lý thuyết thì chỉ còn là tập đề.)",
           italic=True, size=11.5, color=COLOR_MUTED, after=4)

    # Ví dụ minh họa: lấy vài câu DỄ nhất, và đây là chỗ DUY NHẤT có lời giải
    # ngay tại chỗ — đúng như mẫu.
    vi_du = [q for q in tu_luyen if _muc_do_cua(q) == MUC_DE][:3]
    if vi_du:
        _p(doc, "Ví dụ minh họa", bold=True, size=12.5, color=COLOR_PRIMARY,
           before=14, after=4)
        for i, q in enumerate(vi_du, 1):
            _p(doc, f"Ví dụ {i}: {getattr(q, 'new_content', '') or ''}",
               italic=True, size=co_chu, before=8, after=2)
            lg = (getattr(q, "solution_method1", "") or "").strip()
            if lg:
                _p(doc, f"Giải: {lg}", size=co_chu, after=2, thut=18)

    # ---------- PHẦN 2 ----------
    _tieu_de_phan(doc, "PHẦN 2. BÀI TẬP TỰ LUYỆN")
    _p(doc, "Tự làm hết phần này rồi mới mở Phần 4 đối chiếu.",
       italic=True, size=11.5, color=COLOR_MUTED, after=6)

    theo_muc = {m: [q for q in tu_luyen if _muc_do_cua(q) == m] for m in THU_TU_MUC}
    stt = 0
    so_hieu: Dict[int, str] = {}      # id(q) -> nhãn, để Phần 4 gọi đúng tên
    for muc in THU_TU_MUC:
        nhom = theo_muc[muc]
        if not nhom:
            continue
        dau, cuoi = stt + 1, stt + len(nhom)
        khoang = f"Câu {dau}" if dau == cuoi else f"Câu {dau}–{cuoi}"
        _p(doc, f"--- MỨC {muc} ({khoang}) ---", bold=True, size=12,
           color=COLOR_MUC_DO, align=WD_ALIGN_PARAGRAPH.CENTER, before=14, after=6)
        for q in nhom:
            stt += 1
            so_hieu[id(q)] = f"Câu {stt}."
            _in_cau(doc, so_hieu[id(q)], q, co_chu)

    if not tu_luyen:
        _p(doc, "(Không có câu nào cho phần tự luyện.)", italic=True,
           size=11.5, color=COLOR_MUTED)

    # ---------- PHẦN 3 ----------
    _tieu_de_phan(doc, "PHẦN 3. CÂU HỎI TỪ ĐỀ THI")
    if tu_de_thi:
        cac_nam = sorted({n for _, n in tu_de_thi})
        _p(doc, f"Tổng hợp các câu đã xuất hiện trong đề thi, "
                f"trích từ tài liệu gốc ({cac_nam[0]}–{cac_nam[-1]}).",
           italic=True, size=11.5, color=COLOR_MUTED, after=6)
        cum_truoc = ""
        t = 0
        for q, nam in tu_de_thi:
            cum = _cum_nam(nam)
            if cum != cum_truoc:
                cum_truoc = cum
                _p(doc, cum, bold=True, size=12.5, color=COLOR_PRIMARY,
                   before=12, after=4)
            t += 1
            so_hieu[id(q)] = f"Câu T{t}."
            _in_cau(doc, f"Câu T{t}. [{nam}]", q, co_chu)
    else:
        _p(doc, "Tài liệu gốc không có câu nào ghi rõ năm hoặc kỳ thi, nên phần "
                "này để trống. Hãy bổ sung câu từ đề thi thật kèm năm — công cụ "
                "cố ý KHÔNG tự gán thẻ năm cho câu không rõ xuất xứ, vì làm vậy "
                "là bịa nguồn.",
           italic=True, size=11.5, color=COLOR_MUTED, after=4)

    # ---------- PHẦN 4 ----------
    doc.add_page_break()
    _tieu_de_phan(doc, "PHẦN 4. HƯỚNG DẪN VÀ ĐÁP ÁN")

    if tu_luyen:
        # Xếp theo ĐÚNG thứ tự đã in ở Phần 2, không theo thứ tự tài liệu gốc:
        # số hiệu được đánh theo nhóm mức độ nên hai thứ tự này khác nhau, để
        # nguyên thì bảng đáp án chạy Câu 1, Câu 3, Câu 2.
        theo_thu_tu = sorted(tu_luyen, key=lambda q: int(so_hieu[id(q)][4:-1]))
        dong = [(so_hieu[id(q)].rstrip("."),
                 (getattr(q, "correct_answer", "") or "").strip().upper()[:1] or "—",
                 _tom_huong_dan(q)) for q in theo_thu_tu]
        _bang_dap_an(doc, f"A. ĐÁP ÁN PHẦN BÀI TẬP TỰ LUYỆN (Câu 1–{len(tu_luyen)})", dong)

    if tu_de_thi:
        dong = [(so_hieu[id(q)].rstrip("."),
                 (getattr(q, "correct_answer", "") or "").strip().upper()[:1] or "—",
                 _tom_huong_dan(q)) for q, _ in tu_de_thi]
        _bang_dap_an(doc, f"B. ĐÁP ÁN PHẦN CÂU HỎI TỪ ĐỀ THI (Câu T1–T{len(tu_de_thi)})", dong)

    return {
        "so_cau": len(ds),
        "so_cau_tu_luyen": len(tu_luyen),
        "so_cau_de_thi": len(tu_de_thi),
        "theo_muc": {m: len(theo_muc[m]) for m in THU_TU_MUC},
    }
