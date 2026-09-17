# -*- coding: utf-8 -*-
"""
Dựng GIÁO ÁN (Kế hoạch bài dạy) theo khung Phụ lục IV Công văn 5512.

Khung bốn hoạt động, mỗi hoạt động đủ bốn mục:
    Hoạt động 1 — Xác định vấn đề / nhiệm vụ học tập (Mở đầu)
    Hoạt động 2 — Hình thành kiến thức mới
    Hoạt động 3 — Luyện tập
    Hoạt động 4 — Vận dụng
        a) Mục tiêu   b) Nội dung   c) Sản phẩm   d) Tổ chức thực hiện

GIỚI HẠN PHẢI NÓI TRƯỚC, KHÔNG GIẤU:
Giáo án KHÔNG sinh ra được đầy đủ từ một tệp bài tập. Hoạt động 1 và 2 bám vào
nội dung bài học trong sách giáo khoa — thứ mà một tệp bài tập không chứa. Ở đây
Hoạt động 3 và 4 được dựng thật từ các bài trong tài liệu; Hoạt động 1 và 2 để
dạng khung có gợi ý, giáo viên điền nốt. Đầu ra ghi rõ chỗ nào là khung để không
ai tưởng nhầm là đã xong.

Giáo án là HỒ SƠ CHUYÊN MÔN nộp cho tổ và trường, nên ở mọi cấp học đều giữ
nghiêm ngặt: không biểu tượng, không khung trang trí, kể cả tiểu học.
"""

from typing import List, Any, Dict

from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

FONT_MAIN = "Times New Roman"
COLOR_PRIMARY = RGBColor(26, 54, 93)
COLOR_MUTED = RGBColor(74, 85, 104)
COLOR_KHUNG = RGBColor(150, 99, 24)   # phần còn là khung, giáo viên phải điền


def _p(doc, text="", *, bold=False, italic=False, size=12, align=None,
       color=None, before=0, after=0, thut=0):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    if thut:
        para.paragraph_format.left_indent = Pt(thut)
    if text:
        r = para.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.name = FONT_MAIN
        r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
    return para


def _cho_dien(doc, goi_y: str):
    """Dòng khung chờ giáo viên điền — tô khác màu để không lẫn với nội dung thật."""
    _p(doc, f"[Giáo viên bổ sung] {goi_y}", italic=True, size=11.5,
       color=COLOR_KHUNG, thut=18, after=3)


def _muc_hoat_dong(doc, so: int, ten: str):
    _p(doc, f"{so}. Hoạt động {so}: {ten}", bold=True, size=12.5,
       color=COLOR_PRIMARY, before=14, after=4)


def _bon_muc(doc, muc_tieu, noi_dung, san_pham, to_chuc):
    """Bốn mục a) b) c) d) bắt buộc của mỗi hoạt động theo khung 5512."""
    for nhan, gia_tri in (("a) Mục tiêu", muc_tieu), ("b) Nội dung", noi_dung),
                          ("c) Sản phẩm", san_pham), ("d) Tổ chức thực hiện", to_chuc)):
        _p(doc, nhan, bold=True, size=12, thut=12, before=6, after=2)
        if callable(gia_tri):
            gia_tri()
        elif gia_tri:
            _p(doc, gia_tri, size=12, thut=18, after=2)



def _to_chuc(doc, chuyen_giao, thuc_hien, bao_cao, ket_luan):
    """
    Bốn bước tổ chức một hoạt động học, theo Ghi chú 3 của Phụ lục IV.

    Đây là phần dễ viết hời hợt nhất. Văn bản quy định rõ bốn bước có tên, nên
    viết đủ bốn, mỗi bước một đoạn — không gộp thành một câu văn xuôi.
    """
    buoc = (
        ("- Chuyển giao nhiệm vụ", chuyen_giao),
        ("- Thực hiện nhiệm vụ", thuc_hien),
        ("- Báo cáo, thảo luận", bao_cao),
        ("- Kết luận, nhận định", ket_luan),
    )
    for nhan, noi_dung in buoc:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Pt(18)
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.space_after = Pt(2)
        r1 = para.add_run(f"{nhan}: ")
        r1.bold = True
        r1.font.name = FONT_MAIN
        r1.font.size = Pt(12)
        r2 = para.add_run(noi_dung)
        r2.font.name = FONT_MAIN
        r2.font.size = Pt(12)


def xuat_giao_an(doc, book, subject: str = "toan", cap_hoc: str = "THPT",
                 ten_bai: str = "", so_tiet: int = 1,
                 options: Dict[str, Any] = None):
    """Dựng giáo án vào tài liệu Word đang mở."""
    opts = options or {}
    ds: List[Any] = list(getattr(book, "questions", None) or [])
    if not ds:
        for ch in getattr(book, "chapters", None) or []:
            ds.extend(ch.questions)

    mon = "Toán" if subject == "toan" else "Vật lí"
    lop = {"TIEU_HOC": "…", "THCS": "…", "THPT": "…"}.get(cap_hoc, "…")
    ten_bai = (ten_bai or getattr(book, "new_title", "") or "").strip() or "…"

    # ----- Đầu trang theo đúng mẫu Phụ lục IV -----
    _p(doc, "Trường: …………………………………", size=12)
    _p(doc, "Tổ: ………………………………………", size=12)
    _p(doc, "Họ và tên giáo viên: ……………………………", size=12, after=12)

    _p(doc, f"TÊN BÀI DẠY: {ten_bai.upper()}", bold=True, size=14,
       color=COLOR_PRIMARY, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    _p(doc, f"Môn học: {mon}; lớp: {lop}", size=12,
       align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, f"Thời gian thực hiện: {so_tiet} tiết", italic=True, size=12,
       align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

    # ----- I. MỤC TIÊU -----
    _p(doc, "I. MỤC TIÊU", bold=True, size=13, color=COLOR_PRIMARY, before=8, after=4)

    _p(doc, "1. Kiến thức", bold=True, size=12, thut=12, after=2)
    _cho_dien(doc, "Nêu yêu cầu cần đạt của bài này theo chương trình GDPT 2018.")

    _p(doc, "2. Năng lực", bold=True, size=12, thut=12, before=6, after=2)
    _p(doc, f"- Năng lực {'tư duy và lập luận toán học' if subject == 'toan' else 'nhận thức vật lí'}: "
            "phân tích được tình huống, lựa chọn và vận dụng đúng kiến thức đã học.",
       size=12, thut=18, after=2)
    _p(doc, "- Năng lực giải quyết vấn đề: nhận ra vấn đề, đề xuất cách giải và "
            "kiểm tra lại kết quả.", size=12, thut=18, after=2)
    _p(doc, "- Năng lực giao tiếp và hợp tác: trình bày được lời giải, thảo luận "
            "nhóm và nhận xét bài của bạn.", size=12, thut=18, after=2)

    _p(doc, "3. Phẩm chất", bold=True, size=12, thut=12, before=6, after=2)
    _p(doc, "- Chăm chỉ: tự giác hoàn thành nhiệm vụ được giao.", size=12, thut=18, after=2)
    _p(doc, "- Trung thực: báo cáo đúng kết quả làm được, không chép bài.",
       size=12, thut=18, after=2)
    _p(doc, "- Trách nhiệm: hoàn thành phần việc của mình trong hoạt động nhóm.",
       size=12, thut=18, after=2)

    # ----- II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU -----
    _p(doc, "II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU", bold=True, size=13,
       color=COLOR_PRIMARY, before=12, after=4)
    _p(doc, "- Giáo viên: kế hoạch bài dạy, phiếu học tập, máy chiếu.",
       size=12, thut=12, after=2)
    _p(doc, "- Học sinh: sách giáo khoa, vở ghi, đồ dùng học tập.",
       size=12, thut=12, after=2)

    # ----- III. TIẾN TRÌNH DẠY HỌC -----
    _p(doc, "III. TIẾN TRÌNH DẠY HỌC", bold=True, size=13, color=COLOR_PRIMARY,
       before=12, after=4)

    # Hoạt động 1 — khung
    _muc_hoat_dong(doc, 1, "Xác định vấn đề/nhiệm vụ học tập/Mở đầu")
    _bon_muc(
        doc,
        "Tạo hứng thú và làm xuất hiện nhu cầu tìm hiểu kiến thức mới của bài học.",
        lambda: _cho_dien(doc, "Nêu tình huống mở đầu gắn với bài trong sách giáo khoa."),
        "Câu trả lời hoặc dự đoán ban đầu của học sinh (chưa cần chính xác).",
        lambda: _to_chuc(
            doc,
            "giáo viên nêu tình huống mở đầu, yêu cầu cả lớp suy nghĩ cá nhân.",
            "học sinh đọc tình huống và suy nghĩ 1–2 phút; giáo viên quan sát, "
            "gợi ý cho học sinh còn lúng túng.",
            "một vài học sinh phát biểu dự đoán; các học sinh khác nhận xét.",
            "giáo viên chốt vấn đề cần giải quyết và dẫn vào bài học mới.",
        ),
    )

    # Hoạt động 2 — khung
    _muc_hoat_dong(doc, 2, "Hình thành kiến thức mới/giải quyết vấn đề/thực thi nhiệm vụ đặt ra từ Hoạt động 1")
    _bon_muc(
        doc,
        "Học sinh phát biểu được khái niệm, quy tắc hoặc công thức của bài học.",
        lambda: _cho_dien(doc, "Trích nội dung lý thuyết tương ứng trong sách giáo khoa."),
        "Học sinh ghi được nội dung kiến thức vào vở và nêu lại bằng lời.",
        lambda: _to_chuc(
            doc,
            "giáo viên giao nhiệm vụ đọc sách giáo khoa kèm hệ thống câu hỏi dẫn dắt.",
            "học sinh đọc, ghi chép và trả lời câu hỏi; giáo viên theo dõi, hỗ trợ "
            "nhóm gặp khó khăn.",
            "đại diện cặp đôi trình bày kết quả; lớp thảo luận, bổ sung.",
            "giáo viên chốt kiến thức trọng tâm, học sinh ghi vào vở.",
        ),
    )

    # Hoạt động 3 — DỰNG THẬT từ tài liệu
    luyen_tap = [q for q in ds if _muc(q) <= 1] or ds[: max(1, len(ds) // 2)]
    _muc_hoat_dong(doc, 3, "Luyện tập")
    _bon_muc(
        doc,
        "Học sinh vận dụng trực tiếp kiến thức vừa học để giải các bài cơ bản.",
        lambda: _in_bai(doc, luyen_tap, "Bài"),
        "Lời giải đúng của các bài luyện tập, trình bày rõ từng bước.",
        lambda: _to_chuc(
            doc,
            "giáo viên giao hệ thống bài luyện tập, nêu rõ thời gian làm bài.",
            "học sinh làm cá nhân; giáo viên quan sát, hướng dẫn học sinh yếu.",
            "học sinh đổi vở kiểm tra chéo, một số em lên bảng trình bày.",
            "giáo viên chữa bài, nhấn mạnh sai lầm thường gặp và cách tránh.",
        ),
    )

    # Hoạt động 4 — DỰNG THẬT từ tài liệu
    van_dung = [q for q in ds if _muc(q) >= 2]
    if not van_dung:
        van_dung = ds[max(1, len(ds) // 2):] or ds[-1:]
    _muc_hoat_dong(doc, 4, "Vận dụng")
    _bon_muc(
        doc,
        "Học sinh vận dụng kiến thức vào tình huống mới hoặc bài có mức độ cao hơn.",
        lambda: _in_bai(doc, van_dung, "Bài"),
        "Bài làm của học sinh, có lập luận và kết luận rõ ràng.",
        lambda: _to_chuc(
            doc,
            "giáo viên giao nhiệm vụ vận dụng, nêu rõ yêu cầu về sản phẩm.",
            "học sinh thực hiện NGOÀI GIỜ HỌC TRÊN LỚP theo đúng Phụ lục IV.",
            "học sinh nộp báo cáo và trình bày ở tiết sau.",
            "giáo viên nhận xét, đánh giá và ghi nhận vào kết quả học tập.",
        ),
    )

    # ----- Phụ lục: đáp án cho giáo viên -----
    if opts.get("solution", True) and ds:
        doc.add_page_break()
        _p(doc, "PHỤ LỤC — ĐÁP ÁN VÀ LỜI GIẢI (dành cho giáo viên)", bold=True,
           size=13, color=COLOR_PRIMARY, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
        for i, q in enumerate(ds, 1):
            _p(doc, f"Bài {i}. {getattr(q, 'new_content', '')}", bold=True,
               size=11.5, before=8, after=2)
            lg = (getattr(q, "solution_method1", "") or "").strip()
            if lg:
                _p(doc, lg, size=11.5, thut=12, after=2)

    return {
        "so_bai_luyen_tap": len(luyen_tap),
        "so_bai_van_dung": len(van_dung),
        "con_khung": 2,   # Hoạt động 1 và 2 vẫn là khung
    }


def _muc(q) -> int:
    lv = (getattr(q, "level", "") or "").strip().lower()
    if "cao" in lv:
        return 3
    if "vận dụng" in lv or "van dung" in lv:
        return 2
    if "hiểu" in lv or "hieu" in lv:
        return 1
    return 0


def _in_bai(doc, ds, nhan="Bài"):
    for i, q in enumerate(ds, 1):
        _p(doc, f"{nhan} {i}. {getattr(q, 'new_content', '')}",
           size=12, thut=18, after=2)
        for o in (getattr(q, "new_options", None) or []):
            _p(doc, f"    {o}", size=11.5, thut=24, after=0)
