"""
Bộ Kiểm Tra Cấu Trúc & Thẩm Định Văn Bản Trước Xuất Bản (Pre-flight Structural Validator & QA Engine)
Tuân thủ chuẩn Nghị định 30/2020/NĐ-CP và Chương trình GDPT 2018 của Bộ Giáo dục và Đào tạo.

NGUYÊN TẮC: Mọi tiêu chí đều phải được kiểm tra bằng dữ liệu thật.
Không có tiêu chí nào được chấm điểm cứng — một bản báo cáo luôn báo "đạt chuẩn"
là bản báo cáo vô giá trị, vì nó che mất đúng những lỗi cần phát hiện.
"""

import re
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import List, Dict, Any, Optional

# Thể thức bắt buộc theo Nghị định 30/2020/NĐ-CP (đơn vị: milimét)
ND30_PAGE_WIDTH_MM = 210
ND30_PAGE_HEIGHT_MM = 297
ND30_MARGINS_MM = {"left": 30, "right": 15, "top": 20, "bottom": 20}
ND30_TOLERANCE_MM = 1.0


@dataclass
class ValidationCheckItem:
    name: str
    category: str  # "Cấu trúc", "Ngữ nghĩa", "Sư phạm", "Thể thức văn bản"
    passed: bool
    score: int     # Thang điểm của tiêu chí
    max_score: int
    detail: str
    suggestion: str = ""


@dataclass
class ValidationReport:
    is_valid: bool
    total_score: int
    max_score: int = 100
    doc_type: str = "THEMATIC_BOOK"
    total_chapters: int = 0
    total_questions: int = 0
    title_evaluated: str = ""
    checks: List[ValidationCheckItem] = field(default_factory=list)
    summary: str = ""
    warnings: List[str] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        if self.total_score >= 85:
            status = "ĐẠT CHUẨN XUẤT BẢN SƯ PHẠM"
        elif self.total_score >= 70:
            status = "ĐẠT - CẦN RÀ SOÁT MỘT SỐ ĐIỂM"
        else:
            status = "CHƯA ĐẠT - CẦN XỬ LÝ TRƯỚC KHI IN"
        return {
            "is_valid": self.is_valid,
            "total_score": self.total_score,
            "max_score": self.max_score,
            "status_text": status,
            "doc_type": self.doc_type,
            "total_chapters": self.total_chapters,
            "total_questions": self.total_questions,
            "title_evaluated": self.title_evaluated,
            "summary": self.summary,
            "warnings": self.warnings,
            "checks": [asdict(c) for c in self.checks]
        }


def _inspect_exported_docx(path: Path) -> Dict[str, Any]:
    """
    Mở lại chính file Word vừa xuất để đo thể thức thực tế,
    thay vì mặc định tin rằng bộ xuất bản đã làm đúng.
    """
    result: Dict[str, Any] = {"readable": False, "issues": [], "measured": {}}
    try:
        import docx
        from docx.shared import Mm
    except Exception as e:
        result["issues"].append(f"Không nạp được thư viện python-docx để kiểm tra: {e}")
        return result

    try:
        doc = docx.Document(str(path))
    except Exception as e:
        result["issues"].append(f"Không mở lại được file Word vừa xuất: {e}")
        return result

    result["readable"] = True
    section = doc.sections[0]

    def to_mm(emu) -> Optional[float]:
        if emu is None:
            return None
        return round(emu / 36000.0, 2)

    measured = {
        "page_width": to_mm(section.page_width),
        "page_height": to_mm(section.page_height),
        "left": to_mm(section.left_margin),
        "right": to_mm(section.right_margin),
        "top": to_mm(section.top_margin),
        "bottom": to_mm(section.bottom_margin),
    }
    result["measured"] = measured

    if measured["page_width"] is None or abs(measured["page_width"] - ND30_PAGE_WIDTH_MM) > ND30_TOLERANCE_MM:
        result["issues"].append(f"Chiều rộng trang {measured['page_width']}mm khác khổ A4 ({ND30_PAGE_WIDTH_MM}mm)")
    if measured["page_height"] is None or abs(measured["page_height"] - ND30_PAGE_HEIGHT_MM) > ND30_TOLERANCE_MM:
        result["issues"].append(f"Chiều cao trang {measured['page_height']}mm khác khổ A4 ({ND30_PAGE_HEIGHT_MM}mm)")

    for side, expected in ND30_MARGINS_MM.items():
        actual = measured.get(side)
        if actual is None or abs(actual - expected) > ND30_TOLERANCE_MM:
            result["issues"].append(f"Lề {side} đo được {actual}mm, yêu cầu {expected}mm")

    # Đếm số đoạn văn có nội dung để phát hiện file rỗng
    non_empty = sum(1 for p in doc.paragraphs if p.text.strip())
    result["paragraph_count"] = non_empty
    result["table_count"] = len(doc.tables)
    if non_empty < 5:
        result["issues"].append(f"File chỉ có {non_empty} đoạn văn có nội dung — nghi ngờ xuất bản lỗi")

    return result


def _option_label(option_text: str) -> str:
    """Lấy nhãn A/B/C/D ở đầu một phương án."""
    m = re.match(r"\s*([A-Da-d])\s*[.)\-:]", option_text or "")
    return m.group(1).upper() if m else ""


def _option_body(option_text: str) -> str:
    """Lấy phần nội dung của phương án, bỏ nhãn, để so sánh trùng lặp."""
    body = re.sub(r"^\s*[A-Da-d]\s*[.)\-:]\s*", "", option_text or "")
    return re.sub(r"\s+", " ", body).strip().lower().rstrip(".")


class PreFlightValidator:
    """Kiểm tra và thẩm định toàn diện cấu trúc tài liệu trước khi xuất bản"""

    @classmethod
    def validate(
        cls,
        book_title: str,
        subtitle: str,
        chapters: list,
        questions: list,
        doc_type: str = "THEMATIC_BOOK",
        subject: str = "toan",
        exported_path: Optional[Path] = None
    ) -> ValidationReport:
        checks: List[ValidationCheckItem] = []
        warnings: List[str] = []

        all_q_list = []
        if chapters:
            for ch in chapters:
                all_q_list.extend(ch.questions)
        else:
            all_q_list = list(questions or [])

        total_questions = len(all_q_list)
        total_chapters = len(chapters or [])

        # ------------------------------------------------------------------
        # TIÊU CHÍ 1 (20đ): PHÂN TÁCH BÀI TOÁN ĐỘC LẬP
        # Phát hiện hiện tượng cả tài liệu bị dồn thành một khối văn bản.
        # ------------------------------------------------------------------
        c1_max = 20
        overlong = []
        for q in all_q_list:
            body = getattr(q, "new_content", "") or getattr(q, "content", "") or ""
            if len(body) > 2500:
                overlong.append(getattr(q, "title", "?"))

        if total_questions >= 2 and not overlong:
            c1_score, c1_passed = 20, True
            c1_detail = f"Đã phân tách {total_questions} bài toán độc lập, không có bài nào bị dồn cục văn bản."
            c1_sug = "Cấu trúc phân mảnh đạt chuẩn."
        elif total_questions >= 2 and overlong:
            c1_score, c1_passed = 12, False
            c1_detail = (
                f"Phân tách được {total_questions} bài, nhưng {len(overlong)} bài có độ dài bất thường "
                f"(trên 2500 ký tự): {', '.join(overlong[:3])}. Nhiều bài có thể đã bị gộp làm một."
            )
            c1_sug = "Kiểm tra lại quy tắc nhận diện 'Câu N' trong tài liệu nguồn."
            warnings.append(f"{len(overlong)} bài toán có độ dài bất thường, nghi bị gộp.")
        else:
            c1_score, c1_passed = 4, False
            c1_detail = f"Chỉ phát hiện {total_questions} bài toán. Toàn bộ tài liệu có nguy cơ bị gộp thành một khối."
            c1_sug = "Cần phân tích lại cấu trúc phân đoạn (Heading / Table / Paragraph) của tệp nguồn."
            warnings.append("Tài liệu gần như không được phân tách thành các bài riêng biệt.")

        checks.append(ValidationCheckItem(
            name="Phân tách bài toán độc lập",
            category="Cấu trúc", passed=c1_passed, score=c1_score, max_score=c1_max,
            detail=c1_detail, suggestion=c1_sug
        ))

        # ------------------------------------------------------------------
        # TIÊU CHÍ 2 (10đ): PHÂN CẤP CHƯƠNG MỤC
        # ------------------------------------------------------------------
        c2_max = 10
        if total_chapters >= 1:
            c2_score, c2_passed = 10, True
            c2_detail = f"Xác lập {total_chapters} chương/chủ đề chuyên biệt, phân chia đề mục mạch lạc."
            c2_sug = "Cấu trúc phân cấp hoàn chỉnh."
        else:
            c2_score, c2_passed = 7, True
            c2_detail = "Tài liệu dạng đơn đề mục, các bài toán sắp xếp theo chuỗi liên tục."
            c2_sug = "Có thể bổ sung phân chương nếu tài liệu có quy mô lớn."

        checks.append(ValidationCheckItem(
            name="Phân cấp chương mục & đề mục",
            category="Cấu trúc", passed=c2_passed, score=c2_score, max_score=c2_max,
            detail=c2_detail, suggestion=c2_sug
        ))

        # ------------------------------------------------------------------
        # TIÊU CHÍ 3 (10đ): TÊN SÁCH ĐỘC BẢN
        # ------------------------------------------------------------------
        c3_max = 10
        title_clean = (book_title or "").strip()
        generic_names = {"tài liệu", "bộ đề thi", "cuốn sách mới", "tài liệu chuyên đề"}
        is_generic = title_clean.lower() in generic_names or len(title_clean) < 10

        if not is_generic:
            c3_score, c3_passed = 10, True
            c3_detail = f"Tên sách độc bản, phản ánh đúng trọng tâm: «{title_clean}»."
            c3_sug = "Tiêu đề mang phong cách xuất bản chuyên nghiệp."
        else:
            c3_score, c3_passed = 5, False
            c3_detail = f"Tên sách «{title_clean}» còn chung chung hoặc quá ngắn."
            c3_sug = "Dùng chức năng AI Namer để đề xuất tiêu đề độc bản hơn."

        checks.append(ValidationCheckItem(
            name="Tên sách độc bản & nhận diện nội dung",
            category="Ngữ nghĩa", passed=c3_passed, score=c3_score, max_score=c3_max,
            detail=c3_detail, suggestion=c3_sug
        ))

        # ------------------------------------------------------------------
        # TIÊU CHÍ 4 (20đ): TỶ LỆ BÀI CÓ LỜI GIẢI
        # ------------------------------------------------------------------
        c4_max = 20
        has_sol = sum(
            1 for q in all_q_list
            if (getattr(q, "solution_method1", "") or getattr(q, "solution", "") or "").strip()
        )
        sol_ratio = has_sol / max(1, total_questions)
        c4_score = int(round(sol_ratio * c4_max))
        c4_passed = sol_ratio >= 0.8

        if c4_passed:
            c4_detail = f"{has_sol}/{total_questions} bài có lời giải chi tiết (đạt {int(sol_ratio * 100)}%)."
            c4_sug = "Đạt chuẩn sư phạm phát triển năng lực."
        else:
            c4_detail = f"Chỉ {has_sol}/{total_questions} bài có lời giải (đạt {int(sol_ratio * 100)}%)."
            c4_sug = "Bổ sung lời giải cho các bài còn thiếu trước khi in."
            warnings.append(f"{total_questions - has_sol} bài chưa có lời giải.")

        checks.append(ValidationCheckItem(
            name="Chuẩn sư phạm GDPT 2018 & Lời giải",
            category="Sư phạm", passed=c4_passed, score=c4_score, max_score=c4_max,
            detail=c4_detail, suggestion=c4_sug
        ))

        # ------------------------------------------------------------------
        # TIÊU CHÍ 5 (20đ): TÍNH TOÀN VẸN CỦA CÂU TRẮC NGHIỆM
        # Đây là tiêu chí quan trọng nhất về mặt học thuật: phương án trùng nhau,
        # thiếu phương án hoặc đáp án không khớp phương án nào đều là lỗi nặng.
        # ------------------------------------------------------------------
        c5_max = 20
        mcq_list = [q for q in all_q_list if len(getattr(q, "new_options", None) or getattr(q, "options", None) or []) > 0]
        dup_opts, wrong_count, orphan_answer, no_answer = [], [], [], []

        for q in mcq_list:
            opts = getattr(q, "new_options", None) or getattr(q, "options", None) or []
            title = getattr(q, "title", "?")

            bodies = [_option_body(o) for o in opts if _option_body(o)]
            if len(bodies) != len(set(bodies)):
                dup_opts.append(title)

            if len(opts) != 4:
                wrong_count.append(f"{title} ({len(opts)} phương án)")

            answer = (getattr(q, "correct_answer", "") or "").strip().upper()
            if not answer:
                no_answer.append(title)
            else:
                labels = {_option_label(o) for o in opts}
                key = answer[0] if answer and answer[0] in "ABCD" else ""
                if key and labels and key not in labels:
                    orphan_answer.append(title)

        problems = []
        if dup_opts:
            problems.append(f"{len(dup_opts)} câu có phương án trùng lặp ({', '.join(dup_opts[:3])})")
        if wrong_count:
            problems.append(f"{len(wrong_count)} câu không đủ 4 phương án ({', '.join(wrong_count[:3])})")
        if orphan_answer:
            problems.append(f"{len(orphan_answer)} câu có đáp án không khớp phương án nào ({', '.join(orphan_answer[:3])})")
        if no_answer:
            problems.append(f"{len(no_answer)} câu chưa xác định đáp án đúng")

        if not mcq_list:
            c5_score, c5_passed = c5_max, True
            c5_detail = "Tài liệu dạng tự luận, không có câu trắc nghiệm cần đối chiếu phương án."
            c5_sug = "Không áp dụng tiêu chí phương án trắc nghiệm."
        elif not problems:
            c5_score, c5_passed = c5_max, True
            c5_detail = f"{len(mcq_list)} câu trắc nghiệm đều đủ 4 phương án, không trùng lặp và có đáp án hợp lệ."
            c5_sug = "Hệ thống phương án toàn vẹn."
        else:
            faulty = len(set(dup_opts) | set(orphan_answer)) + len(wrong_count) + len(no_answer)
            ratio_ok = max(0.0, 1 - faulty / max(1, len(mcq_list)))
            c5_score = int(round(ratio_ok * c5_max))
            c5_passed = False
            c5_detail = "Phát hiện lỗi phương án trắc nghiệm: " + "; ".join(problems) + "."
            c5_sug = "Rà soát lại các câu nêu trên — phương án trùng hoặc thiếu đáp án sẽ gây sai lệch khi dạy học."
            warnings.extend(problems)

        checks.append(ValidationCheckItem(
            name="Tính toàn vẹn phương án & đáp án trắc nghiệm",
            category="Sư phạm", passed=c5_passed, score=c5_score, max_score=c5_max,
            detail=c5_detail, suggestion=c5_sug
        ))

        # ------------------------------------------------------------------
        # TIÊU CHÍ 6 (20đ): THỂ THỨC VĂN BẢN NGHỊ ĐỊNH 30/2020/NĐ-CP
        # Đo trực tiếp trên file Word đã xuất, không chấm điểm cứng.
        # ------------------------------------------------------------------
        c6_max = 20
        if exported_path is None:
            c6_score, c6_passed = 10, False
            c6_detail = "Chưa có file Word thành phẩm để đo thể thức — tiêu chí này chỉ được chấm sau khi xuất bản."
            c6_sug = "Gọi hàm thẩm định sau bước xuất file để đo lề và khổ giấy thực tế."
        else:
            info = _inspect_exported_docx(Path(exported_path))
            m = info.get("measured", {})
            if not info.get("readable"):
                c6_score, c6_passed = 0, False
                c6_detail = "Không mở lại được file Word vừa xuất: " + "; ".join(info.get("issues", []))
                c6_sug = "Kiểm tra lại bộ xuất bản DocxBookExporter."
                warnings.append("File Word thành phẩm không đọc lại được.")
            elif info.get("issues"):
                c6_score, c6_passed = max(0, c6_max - 5 * len(info["issues"])), False
                c6_detail = "Thể thức chưa đạt: " + "; ".join(info["issues"])
                c6_sug = "Điều chỉnh lại thiết lập trang trong DocxBookExporter.apply_standard_page_setup."
                warnings.extend(info["issues"])
            else:
                c6_score, c6_passed = c6_max, True
                c6_detail = (
                    f"Đo trực tiếp trên file thành phẩm: khổ {m.get('page_width')}x{m.get('page_height')}mm (A4); "
                    f"lề Trái {m.get('left')}mm, Phải {m.get('right')}mm, "
                    f"Trên {m.get('top')}mm, Dưới {m.get('bottom')}mm; "
                    f"{info.get('paragraph_count', 0)} đoạn văn, {info.get('table_count', 0)} bảng."
                )
                c6_sug = "Thể thức đúng chuẩn Nghị định 30/2020/NĐ-CP."

        checks.append(ValidationCheckItem(
            name="Thể thức văn bản Nghị định 30/2020/NĐ-CP",
            category="Thể thức văn bản", passed=c6_passed, score=c6_score, max_score=c6_max,
            detail=c6_detail, suggestion=c6_sug
        ))

        total_score = sum(c.score for c in checks)
        is_valid = total_score >= 85 and not warnings

        summary = (
            f"Đạt {total_score}/100 điểm thẩm định. "
            f"Tài liệu gồm {total_chapters} chương/chủ đề và {total_questions} bài toán."
        )
        if warnings:
            summary += f" Có {len(warnings)} điểm cần rà soát trước khi in."

        return ValidationReport(
            is_valid=is_valid,
            total_score=total_score,
            max_score=100,
            doc_type=doc_type,
            total_chapters=total_chapters,
            total_questions=total_questions,
            title_evaluated=book_title,
            checks=checks,
            summary=summary,
            warnings=warnings
        )
